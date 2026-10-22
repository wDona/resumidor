"""Exportar el documento (docx, pdf, md, obsidian, notion, ics, csv, srt/vtt), editarlo conservando el original y estilos de
organización. Todo sale de los datos estructurados (acta.json), nunca del HTML. Edición: acta_editada.json junto al original."""
import copy
import csv
import datetime as dt
import html
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

import pipeline as P
import plantillas

EDITADA = "acta_editada.json"
CAMPOS = [("numero", "Nº"), ("tipo", "Tipo de sesión"), ("lugar", "Lugar"), ("hora_inicio", "Hora de inicio"), ("hora_fin", "Hora de fin"),
          ("convocatoria", "Convocatoria"), ("asistentes", "Asistentes"), ("ausentes", "Ausentes")]
FORMATOS = {"docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "pdf": ("pdf", "application/pdf"),
            "md": ("md", "text/markdown; charset=utf-8"), "obsidian": ("obsidian.zip", "application/zip"), "notion": ("notion.zip", "application/zip"),
            "ics": ("ics", "text/calendar; charset=utf-8"), "csv_todoist": ("todoist.csv", "text/csv; charset=utf-8"),
            "csv_jira": ("jira.csv", "text/csv; charset=utf-8"), "srt": ("srt", "application/x-subrip"), "vtt": ("vtt", "text/vtt; charset=utf-8")}


class SinNavegador(Exception):
    pass


# ---------- carga ----------
def carga(mid, ver, anon=False, dudas=False):
    """anon: la copia anonimizada (privacidad.aplicar) en vez del documento; dudas: marca lo dudoso (dudas.py) en los exports."""
    d = P.WORK / Path(mid).name / Path(ver).name
    if not (d / "acta.json").is_file():
        raise FileNotFoundError(ver)
    ed = d / EDITADA
    if anon and not (d / "acta_anonima.json").is_file():
        raise ValueError("No hay copia anonimizada de este documento: anonimízalo primero")
    data = json.loads((d / "acta_anonima.json" if anon else ed if ed.is_file() else d / "acta.json").read_text())
    tf = d.parent / f"transcript_{d.name.split('_')[2]}.json"
    try:
        tpl = next(t for t in plantillas.todas() if t["nombre"] == data.get("plantilla")) if data.get("plantilla") else plantillas.get("acta")
    except StopIteration:
        tpl = plantillas.get("acta")
    try:
        carpeta = json.loads((d.parent / "meta.json").read_text()).get("carpeta", "")
    except (OSError, ValueError):
        carpeta = ""
    if anon:
        tf = d / "transcript_anonimo.json"
    c = {"d": d, "mid": d.parent.name, "ver": d.name, "data": data, "tpl": tpl, "editado": ed.is_file() and not anon, "carpeta": carpeta, "anon": anon,
         "img": "img_anon" if anon else "img",
         "segs": json.loads(tf.read_text()).get("segments", []) if tf.is_file() else [],
            "doc": (data.get("documento") or tpl.get("documento") or {}).get("singular", "acta"),
         "fecha": dt.datetime.fromtimestamp((d / "acta.json").stat().st_mtime)}
    if dudas:
        import dudas as DU
        c["dudas"] = DU.calcular(mid, ver)["dudas"]
    return c


# ---------- edición con copia ----------
RUTA = re.compile(r"[a-z_]+(\.\d+|\.[a-z_]+)*")


def _set(data, ruta, texto):
    if not RUTA.fullmatch(ruta):
        raise ValueError(f"Ruta no válida: {ruta}")
    *ks, last = [int(k) if k.isdigit() else k for k in ruta.split(".")]
    o = data
    try:
        for k in ks:
            o = o[k]
        if not isinstance(o[last], str):
            raise ValueError(f"No es texto: {ruta}")
        o[last] = str(texto).strip()
    except (KeyError, IndexError, TypeError):
        raise ValueError(f"No existe: {ruta}")


def guardar_edicion(mid, ver, cambios):
    """Aplica {ruta: texto} sobre la versión vigente (editada u original) y escribe acta_editada.json; el original no se toca."""
    c = carga(mid, ver)
    data = copy.deepcopy(c["data"])
    for r, t in cambios.items():
        _set(data, r, t)
    tmp = c["d"] / (EDITADA + ".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(c["d"] / EDITADA)


def restaurar(mid, ver):
    (P.WORK / Path(mid).name / Path(ver).name / EDITADA).unlink(missing_ok=True)


# ---------- bloques (orden y títulos del perfil) ----------
def _sub(x):
    return " · ".join(f"{k.replace('_', ' ').capitalize()}: {v}" for k, v in x.items()
                      if k not in ("cita", "texto", "importancia", "t") and isinstance(v, str) and v.strip(" —-"))


def bloques(c):
    """[(clase, título, items)]: clase texto|parrafos|lista|tareas|secciones|capturas|trans. Cada item lleva `k` (ruta editable) y `t`."""
    D, T = c["data"], {**P.ACTA_TIT, **c["tpl"].get("titulos", {})}
    sps = {s["id"]: (i, s) for i, s in enumerate(D.get("secciones_plantilla", []))}
    out = []
    for b in c["tpl"].get("orden") or P.ACTA_ORDEN:
        tit = T.get(b) or (sps[b][1]["nombre"] if b in sps else b.capitalize())
        if b == "resumen" and D.get("resumen"):
            out.append(("texto", tit, [{"k": "resumen", "texto": D["resumen"]}]))
        elif b == "desarrollo" and D.get("relato"):
            out.append(("parrafos", tit, [{"k": f"relato.{i}.texto", "t": x["t"], "texto": x["texto"]} for i, x in enumerate(D["relato"])]))
        elif b == "decisiones" and D.get("decisiones"):
            out.append(("lista", tit, [{"k": f"decisiones.{i}.texto", "t": x["t"], "texto": x["texto"], "cita": x.get("cita", "")} for i, x in enumerate(D["decisiones"])]))
        elif b == "tareas" and D.get("tareas"):
            out.append(("tareas", tit, [{"k": f"tareas.{i}.tarea", "kr": f"tareas.{i}.responsable", "kp": f"tareas.{i}.plazo", "t": x["t"], "texto": x["tarea"],
                                         "resp": x.get("responsable", ""), "plazo": x.get("plazo", "")} for i, x in enumerate(D["tareas"])]))
        elif b == "detalle":
            pts = [{"k": f"puntos.{i}.texto", "t": x["t"], "texto": x["texto"]} for i, x in enumerate(D.get("puntos", []))]
            ss = D.get("secciones", [])
            if ss:
                out.append(("secciones", tit, [{"t": s["t"], "titulo": s["titulo"], "kt": f"secciones.{i}.titulo", "resumen": s["resumen"], "kr": f"secciones.{i}.resumen",
                                                "puntos": [p for p in pts if (i == 0 or p["t"] >= s["t"]) and (i + 1 == len(ss) or p["t"] < ss[i + 1]["t"])]}
                                               for i, s in enumerate(ss)]))
            elif pts:
                out.append(("lista", "Puntos clave", pts))
        elif b in sps and sps[b][1]["items"]:
            i, s = sps[b]
            out.append(("lista", tit, [{"k": f"secciones_plantilla.{i}.items.{j}.texto", "t": x["t"], "texto": x["texto"], "sub": _sub(x), "cita": x.get("cita", "")}
                                       for j, x in enumerate(s["items"])]))
        elif b == "transcripcion" and c["segs"]:
            out.append(("capturas", "Capturas", [k for k in D.get("capturas", []) if (c["d"] / c["img"] / k["img"]).is_file()]))
            out.append(("trans", tit, c["segs"]))
    if not any(k == "capturas" for k, _, _ in out):  # perfil sin transcripción en el orden: las capturas van al final
        out.append(("capturas", "Capturas", [k for k in D.get("capturas", []) if (c["d"] / c["img"] / k["img"]).is_file()]))
    out = [b for b in out if b[2]]
    if c.get("dudas"):
        _marcar(out, c["dudas"])
    return out


def _marcar(out, dud):
    """&dudas=1: lo dudoso lleva «⚠ » delante (md, docx, pdf; en HTML además va en <mark> con los motivos)."""
    def m(x, kc, campo):
        if x.get(kc) in dud and x.get(campo):
            x[campo] = "⚠ " + x[campo]
    for cl, _, its in out:
        for x in its if cl in ("texto", "parrafos", "lista", "tareas", "secciones") else []:
            m(x, "k", "texto")
            m(x, "kt", "titulo")
            if cl == "tareas":
                m(x, "kr", "resp")
                m(x, "kp", "plazo")
            if cl == "secciones":
                m(x, "kr", "resumen")
                for p in x["puntos"]:
                    m(p, "k", "texto")


fm = P.fmt
esc = html.escape


# ---------- HTML (vista editable, impresión y PDF) ----------
CSS = """:root{--bg:#f4f5f7;--card:#fff;--ink:#1f2933;--mut:#616e7c;--line:#d9dee4;--line2:#e6e9ee;--acc:#1d5fd1;--acc-d:#174ba3;--acc-s:#eef4ff;color-scheme:light dark}
@media(prefers-color-scheme:dark){:root{--bg:#14171c;--card:#1d222a;--ink:#e6e9ee;--mut:#9aa5b1;--line:#333b46;--line2:#2a313a;--acc:#6ea3ff;--acc-d:#8db4f5;--acc-s:#1b2a45}}
*{box-sizing:border-box}body{font:16px/1.6 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif;max-width:960px;margin:0 auto;padding:32px 40px;color:var(--ink);background:var(--card)}
@media(max-width:700px){body{padding:20px 16px}}
h1{font-size:26px;line-height:1.25;margin:0 0 6px}h2{font-size:22px;line-height:1.3;margin:1.8em 0 .6em}h3{font-size:17px;margin:1.4em 0 .4em}
.ts{display:inline-block;font:600 13px system-ui,sans-serif;font-variant-numeric:tabular-nums;background:var(--acc-s);color:var(--acc-d);padding:1px 7px;border-radius:5px;margin-right:6px}.ts[data-t]{cursor:pointer}
@media(prefers-color-scheme:dark){.ts{color:var(--acc)}}
.meta,small{color:var(--mut);font-size:14px}.cab{display:flex;align-items:center;gap:12px;border-bottom:3px solid var(--acc);padding-bottom:10px;margin-bottom:18px;color:var(--mut)}.cab img{max-height:48px}
table{border-collapse:collapse;width:100%;font-size:15px}td,th{border-top:1px solid var(--line2);padding:8px 10px;text-align:left;vertical-align:top}th{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--mut);border-top:0}
figure{margin:12px 0 16px}img{max-width:100%;border-radius:10px}figcaption{font-size:13px;color:var(--mut)}
[contenteditable]{outline:1px dashed var(--line);border-radius:3px}[contenteditable]:focus{outline:2px solid var(--acc)}:focus-visible{outline:3px solid #8db4f5;outline-offset:2px}
mark{background:#ffe58a;color:#000;border-radius:3px;padding:0 2px}@media print{body{max-width:none;padding:0;color:#000;background:#fff}figure,tr{break-inside:avoid}.ts{background:none;padding:0;color:#555}}"""
JS_EDIT = ("document.querySelectorAll('[data-p]').forEach(e=>e.dataset.o=e.innerText.trim());"
           "window.cambios=()=>{const o={};document.querySelectorAll('[data-p]').forEach(e=>{const t=e.innerText.trim();if(t!==e.dataset.o)o[e.dataset.p]=t});return o}")


def cabecera_estilo(est, base=None):
    """Franja con logo y cabecera del estilo de organización (solo en la vista de la app, donde el logo se sirve por la API)."""
    if not est or not (est.get("cabecera") or est.get("logo")):
        return ""
    logo = f'<img src="/api/estilos/{esc(est["id"])}/logo" alt="">' if est.get("logo") and (base or "").startswith("/") else ""
    return f'<div class="cab">{logo}<span>{esc(est.get("cabecera", ""))}</span></div>'


def html_doc(c, editar=False, base=None, imprimir=False, estilo=None):
    """base: URL de la que cuelgan las imágenes (img/…). estilo: color y cabecera/logo del estilo de organización."""
    dud = c.get("dudas") or {}
    ed = lambda k, t: f'<mark data-p="{k}" title="{esc("; ".join(dud[k]))}">{esc(t)}</mark>' if k in dud else f'<span data-p="{k}"{" contenteditable" if editar else ""}>{esc(t)}</span>'
    ts = lambda t: f'<span class="ts" data-t="{int(t)}">[{fm(t)}]</span>'
    D, h = c["data"], []
    for cl, tit, its in bloques(c):
        h.append(f"<h2>{esc(tit)}</h2>")
        if cl == "texto":
            h.append(f"<p>{ed(its[0]['k'], its[0]['texto'])}</p>")
        elif cl == "parrafos":
            h += [f"<p>{ts(x['t'])}{ed(x['k'], x['texto'])}</p>" for x in its]
        elif cl == "lista":
            h.append("<ul>" + "".join(f"<li>{ts(x['t'])}{ed(x['k'], x['texto'])}" + (f" <small>— {esc(x['sub'])}</small>" if x.get("sub") else "") + "</li>" for x in its) + "</ul>")
        elif cl == "tareas":
            h.append("<table><tr><th>Tarea</th><th>Responsable</th><th>Plazo</th></tr>" + "".join(
                f"<tr><td>{ts(x['t'])}{ed(x['k'], x['texto'])}</td><td>{ed(x['kr'], x['resp'])}</td><td>{ed(x['kp'], x['plazo'])}</td></tr>" for x in its) + "</table>")
        elif cl == "secciones":
            for s in its:
                h.append(f"<h3>{ts(s['t'])}{ed(s['kt'], s['titulo'])}</h3><p>{ed(s['kr'], s['resumen'])}</p><ul>" +
                         "".join(f"<li>{ts(x['t'])}{ed(x['k'], x['texto'])}</li>" for x in s["puntos"]) + "</ul>")
        elif cl == "capturas":
            h += [f'<figure><img src="{c["img"]}/{esc(k["img"])}" loading="lazy"><figcaption>{ts(k["t"])}{esc(k.get("titulo") or "Pantalla compartida")}</figcaption></figure>' for k in its]
        elif cl == "trans":
            h += [f"<p>{ts(s['s'])}" + (f"<b>{esc(s['spk'])}:</b> " if s.get("spk") else "") + f"{esc(s['text'])}</p>" for s in its]
    cab = c["doc"].capitalize()
    return (f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{esc(cab)} — {esc(D["titulo"])}</title>'
            + (f'<base href="{esc(base)}">' if base else "") + f"<style>{CSS}{(':root{--acc:' + estilo['color'] + ';--acc-d:' + estilo['color'] + '}') if estilo and estilo.get('color') else ''}</style></head><body>{cabecera_estilo(estilo, base)}<h1>{esc(cab)}: {esc(D['titulo'])}</h1>"
            f'<p class="meta">Duración {fm(D["duracion"])} · {c["fecha"]:%Y-%m-%d}' + (" · editado" if c["editado"] else "") + "</p>" + "".join(h)
            + (f"<script>{JS_EDIT}</script>" if editar else "") + ("<script>onload=()=>print()</script>" if imprimir else "") + "</body></html>")


# ---------- PDF con el navegador instalado ----------
def navegador():
    nombres = ["msedge", "microsoft-edge", "microsoft-edge-stable", "google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome", "brave-browser", "brave"]
    ruta = [r"%ProgramFiles(x86)%\Microsoft\Edge\Application\msedge.exe", r"%ProgramFiles%\Microsoft\Edge\Application\msedge.exe",
            r"%ProgramFiles%\Google\Chrome\Application\chrome.exe", r"%ProgramFiles(x86)%\Google\Chrome\Application\chrome.exe",
            r"%LocalAppData%\Google\Chrome\Application\chrome.exe", r"%ProgramFiles%\BraveSoftware\Brave-Browser\Application\brave.exe",
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
            "/Applications/Chromium.app/Contents/MacOS/Chromium", "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser"]
    return next(filter(None, [shutil.which(n) for n in nombres] + [p for p in map(os.path.expandvars, ruta) if os.path.isfile(p)]), None)


def pdf_bytes(c, estilo=None):
    exe = navegador()
    if not exe:
        raise SinNavegador("Usa Imprimir → Guardar como PDF")
    with tempfile.TemporaryDirectory() as t:
        f, out = Path(t) / "doc.html", Path(t) / "doc.pdf"
        f.write_text(html_doc(c, base=c["d"].as_uri() + "/", estilo=estilo), encoding="utf-8")
        cmd = [exe, "--headless=new", "--disable-gpu", "--no-sandbox", f"--user-data-dir={t}/perfil", "--no-pdf-header-footer", f"--print-to-pdf={out}", f.as_uri()]
        pr = subprocess.Popen(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)  # sin tuberías: los hijos del navegador las retienen
        fin, last = time.time() + 90, -1
        while time.time() < fin and pr.poll() is None:
            time.sleep(0.3)
            sz = out.stat().st_size if out.is_file() else 0
            if sz and sz == last:  # el PDF ya está escrito y estable
                break
            last = sz
        if pr.poll() is None:
            pr.terminate()
            try:
                pr.wait(5)
            except subprocess.TimeoutExpired:
                pr.kill()
        if not out.is_file() or not out.stat().st_size:
            raise SinNavegador("No se pudo generar el PDF con el navegador. Usa Imprimir → Guardar como PDF")
        return out.read_bytes()


# ---------- Markdown / Obsidian / Notion ----------
def _fila(pers, x):
    return f"{x['texto']}" + (f" ({x['sub']})" if x.get("sub") else "")


def md(c, img=None, links=False, plano=True):
    """img: carpeta relativa de las capturas (None = sin imágenes). links: [[Persona]] en responsables y hablantes."""
    pers = (lambda n: f"[[{n}]]" if links and n else n)
    ts = lambda t: f"[{fm(t)}] "
    o = [f"# {c['doc'].capitalize()}: {c['data']['titulo']}\n"]
    for cl, tit, its in bloques(c):
        o.append(f"\n## {tit}\n")
        if cl == "texto":
            o.append(its[0]["texto"])
        elif cl == "parrafos":
            o += [f"{ts(x['t'])}{x['texto']}\n" for x in its]
        elif cl == "lista":
            o += [f"- {ts(x['t'])}{_fila(pers, x)}" for x in its]
        elif cl == "tareas":
            o += [f"- [ ] {x['texto']}" + (f" — {pers(x['resp'])}" if x["resp"].strip(" —-") else "") + (f" ({x['plazo']})" if x["plazo"].strip(" —-") else "") + f" {ts(x['t']).strip()}" for x in its]
        elif cl == "secciones":
            for s in its:
                o.append(f"\n### {ts(s['t'])}{s['titulo']}\n\n{s['resumen']}\n")
                o += [f"- {ts(x['t'])}{x['texto']}" for x in s["puntos"]]
        elif cl == "capturas" and img is not None:
            o += [f"![{ts(k['t'])}{k.get('titulo') or 'Pantalla compartida'}]({img}/{k['img']})\n" for k in its]
        elif cl == "trans":
            o += [f"{ts(s['s'])}" + (f"**{pers(s['spk'])}:** " if s.get("spk") else "") + s["text"] + "\n" for s in its]
    return "\n".join(o) + "\n"


def _nombre_arch(s):
    return re.sub(r'[\\/:*?"<>|#^\[\]\n]+', " ", s).strip()[:80] or "documento"


def frontmatter(c):
    pr = sorted({x["resp"] for cl, _, its in bloques(c) if cl == "tareas" for x in its if x["resp"].strip(" —-")})
    return "\n".join(["---", f"title: {json.dumps(c['data']['titulo'], ensure_ascii=False)}", f"date: {c['fecha']:%Y-%m-%d}", f"tipo: {c['doc']}",
                      f"duracion: {fm(c['data']['duracion'])}"] + ([f"carpeta: \"[[{c['carpeta']}]]\""] if c["carpeta"] else [])
                     + (["personas:"] + [f"  - \"[[{p}]]\"" for p in pr] if pr else []) + ["tags: [actas]", "---", ""])


def obsidian_zip(c):
    nom, b = _nombre_arch(f"{c['data']['titulo']}"), io.BytesIO()
    pers = {x["resp"] for cl, _, its in bloques(c) if cl == "tareas" for x in its if x["resp"].strip(" —-")} | {s["spk"] for s in c["segs"] if s.get("spk")}
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{nom}.md", frontmatter(c) + md(c, img="adjuntos", links=True))
        for p in sorted(pers):
            z.writestr(f"Personas/{_nombre_arch(p)}.md", f"---\ntags: [persona]\n---\n# {p}\n\nAparece en [[{nom}]].\n")
        if c["carpeta"]:
            z.writestr(f"Carpetas/{_nombre_arch(c['carpeta'])}.md", f"---\ntags: [carpeta]\n---\n# {c['carpeta']}\n\n- [[{nom}]]\n")
        for k in c["data"].get("capturas", []):
            if (c["d"] / c["img"] / k["img"]).is_file():
                z.write(c["d"] / c["img"] / k["img"], f"adjuntos/{k['img']}")
    return b.getvalue()


def notion_zip(c):
    b, nom = io.BytesIO(), _nombre_arch(c["data"]["titulo"])
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"{nom}.md", md(c, img="img"))
        for k in c["data"].get("capturas", []):
            if (c["d"] / c["img"] / k["img"]).is_file():
                z.write(c["d"] / c["img"] / k["img"], f"img/{k['img']}")
    return b.getvalue()


# ---------- fechas, ICS y CSV ----------
MESES = {m: i + 1 for i, m in enumerate("enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre".split())} | {"setiembre": 9}


def fecha(txt, ref):
    """Plazo en texto -> date o None: 2026-03-15 · 15/03/2026 · 15 de marzo (de 2026); sin año, el del documento."""
    s = (txt or "").lower()
    try:
        if m := re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", s):
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        if m := re.search(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})\b", s):
            return dt.date(int(m[3]) + (2000 if len(m[3]) == 2 else 0), int(m[2]), int(m[1]))
        if (m := re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)(?:\s+de(?:l)?\s+(\d{4}))?", s)) and m[2] in MESES:
            return dt.date(int(m[3] or ref.year), MESES[m[2]], int(m[1]))
    except ValueError:
        pass
    return None


def tareas(c):
    return [x for cl, _, its in bloques(c) if cl == "tareas" for x in its]


def _desc(c, x):
    return (f"Responsable: {x['resp']}\n" if x["resp"].strip(" —-") else "") + (f"Plazo: {x['plazo']}\n" if x["plazo"].strip(" —-") else "") + f"Minuto {fm(x['t'])} de «{c['data']['titulo']}»"


def ics(c, evento=False):
    esc_ = lambda s: s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")

    def fold(l):
        out, cur = [], ""
        for ch in l:
            if len((cur + ch).encode()) > (74 if out else 75):
                out.append(cur)
                cur = ch
            else:
                cur += ch
        return "\r\n ".join(out + [cur])
    now, L = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%dT%H%M%SZ"), ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//actas//ES", "CALSCALE:GREGORIAN"]
    base = f"{c['mid']}-{c['ver']}".replace(" ", "_")
    if evento:
        a = c["fecha"]
        L += ["BEGIN:VEVENT", f"UID:{base}-evento@actas", f"DTSTAMP:{now}", f"DTSTART:{a:%Y%m%dT%H%M%S}", f"DTEND:{a + dt.timedelta(seconds=c['data']['duracion']):%Y%m%dT%H%M%S}",
              f"SUMMARY:{esc_(c['data']['titulo'])}", f"DESCRIPTION:{esc_(c['data'].get('resumen', ''))}", "END:VEVENT"]
    for i, x in enumerate(tareas(c)):
        f = fecha(x["plazo"], c["fecha"])
        L += ["BEGIN:VTODO", f"UID:{base}-{i}@actas", f"DTSTAMP:{now}", f"SUMMARY:{esc_(x['texto'])}", f"DESCRIPTION:{esc_(_desc(c, x))}"]
        L += ([f"DUE;VALUE=DATE:{f:%Y%m%d}"] if f else []) + ([f"CATEGORIES:{esc_(x['resp'])}"] if x["resp"].strip(" —-") else []) + ["STATUS:NEEDS-ACTION", "END:VTODO"]
    return "\r\n".join(fold(l) for l in L + ["END:VCALENDAR", ""])


def _csv(cab, filas):
    b = io.StringIO()
    w = csv.writer(b, lineterminator="\r\n")
    w.writerow(cab)
    w.writerows(filas)
    return "\ufeff" + b.getvalue()


def csv_todoist(c):
    return _csv(["TYPE", "CONTENT", "DESCRIPTION", "PRIORITY", "INDENT", "AUTHOR", "RESPONSIBLE", "DATE", "DATE_LANG", "TIMEZONE"],
                [["task", x["texto"], _desc(c, x), 4, 1, "", x["resp"].strip(" —-"), (f"{f:%Y-%m-%d}" if (f := fecha(x["plazo"], c["fecha"])) else x["plazo"].strip(" —-")), "es", ""]
                 for x in tareas(c)])


def csv_jira(c):
    return _csv(["Summary", "Issue Type", "Description", "Assignee", "Due date"],
                [[x["texto"], "Task", _desc(c, x), x["resp"].strip(" —-"), f"{f:%Y-%m-%d}" if (f := fecha(x["plazo"], c["fecha"])) else ""] for x in tareas(c)])


# ---------- subtítulos ----------
def palabras(s):
    """[(palabra, ini, fin)]: con tiempos de palabra si el segmento los tiene; si no, repartido en proporción a los caracteres."""
    if s.get("words"):
        return [(w["w"], w["s"], w["e"]) for w in s["words"] if w.get("w")]
    tk = s["text"].split()
    n, cum, out = sum(len(t) + 1 for t in tk) or 1, 0, []
    for t in tk:
        out.append((t, s["s"] + cum / n * (s["e"] - s["s"]), s["s"] + (cum + len(t)) / n * (s["e"] - s["s"])))
        cum += len(t) + 1
    return out


def cues(segs, hablantes=False, ancho=42, lineas=2):
    out = []
    for i, s in enumerate(segs):
        ws = palabras(s)
        if not ws:
            continue
        if hablantes and s.get("spk") and (i == 0 or segs[i - 1].get("spk") != s["spk"]):
            ws = [(s["spk"] + ":", ws[0][1], ws[0][1])] + ws
        ls, a, b = [""], None, None

        def vaciar():
            nonlocal ls, a, b
            if a is not None:
                out.append([a, b, "\n".join(ls)])
            ls, a, b = [""], None, None
        for w, ini, fin in ws:
            for p in [w[k:k + ancho] for k in range(0, len(w), ancho)]:
                if ls[-1] and len(ls[-1]) + 1 + len(p) > ancho:
                    if len(ls) == lineas:
                        vaciar()
                    else:
                        ls.append("")
                ls[-1] = (ls[-1] + " " + p).strip()
                a = ini if a is None else a
                b = fin
        vaciar()
    prev = 0.0
    for q in out:  # tiempos crecientes y sin solaparse
        q[0] = max(q[0], prev)
        q[1] = max(q[1], q[0] + 0.2)
        prev = q[1]
    return out


def subtitulos(segs, vtt=False, hablantes=False):
    f = lambda t: f"{int(t // 3600):02d}:{int(t % 3600 // 60):02d}:{int(t % 60):02d}{'.' if vtt else ','}{int(round(t % 1 * 1000)) % 1000:03d}"
    return ("WEBVTT\n\n" if vtt else "") + "\n".join(f"{'' if vtt else str(i) + chr(10)}{f(a)} --> {f(b)}\n{t}\n" for i, (a, b, t) in enumerate(cues(segs, hablantes), 1))


# ---------- estilos de organización ----------
def _est_dir():
    d = P.WORK / "estilos"
    d.mkdir(parents=True, exist_ok=True)
    return d


ID = re.compile(r"[a-z0-9][a-z0-9-]{0,39}")


def estilos():
    return [json.loads(f.read_text()) for f in sorted(_est_dir().glob("*.json"))]


def estilo_get(eid):
    f = _est_dir() / f"{eid}.json"
    if not ID.fullmatch(eid or "") or not f.is_file():
        raise KeyError(eid)
    return json.loads(f.read_text())


def estilo_guardar(d):
    """Crea (sin id) o actualiza un estilo; logo y base_docx solo se conservan o se quitan ("" ) aquí: se suben aparte."""
    nombre = (d.get("nombre") or "").strip()[:80]
    if not nombre:
        raise ValueError("Falta el nombre")
    eid = d.get("id")
    if not eid:
        s = re.sub(r"[^a-z0-9]+", "-", nombre.lower().translate(str.maketrans("áéíóúüñç", "aeiouunc"))).strip("-")[:30] or "estilo"
        eid, i = s, 1
        while (_est_dir() / f"{eid}.json").exists():
            i += 1
            eid = f"{s}-{i}"
    elif not ID.fullmatch(eid):
        raise ValueError("Id no válido")
    prev = estilo_get(eid) if (_est_dir() / f"{eid}.json").exists() else {}
    color = (d.get("color") or "").strip()
    if color and not re.fullmatch(r"#[0-9a-fA-F]{6}", color):
        raise ValueError("Color: #rrggbb")
    e = {"id": eid, "nombre": nombre, "logo": prev.get("logo", ""), "base_docx": prev.get("base_docx", ""), "cabecera": str(d.get("cabecera", ""))[:300],
         "pie": str(d.get("pie", ""))[:300], "color": color, "campos": {k: str((d.get("campos") or {}).get(k, ""))[:500] for k, _ in CAMPOS},
         "firmas": [{"cargo": str(f.get("cargo", ""))[:80], "nombre": str(f.get("nombre", ""))[:80]} for f in (d.get("firmas") or [])[:8]]}
    for k in ("logo", "base_docx"):
        if k in d and not d[k] and e[k]:
            (_est_dir() / e[k]).unlink(missing_ok=True)
            e[k] = ""
    (_est_dir() / f"{eid}.json").write_text(json.dumps(e, ensure_ascii=False, indent=1))
    return e


def estilo_fichero(eid, clave, nombre, datos):
    """Guarda el logo (png/jpg) o el .docx base de un estilo existente."""
    e = estilo_get(eid)
    if clave == "logo":
        ext = Path(nombre).suffix.lower()
        if ext not in (".png", ".jpg", ".jpeg") or len(datos) > 5_000_000:
            raise ValueError("Logo: png o jpg de hasta 5 MB")
        from PIL import Image
        Image.open(io.BytesIO(datos)).verify()
        fn = f"{eid}_logo{ext}"
    else:
        from docx import Document
        if len(datos) > 20_000_000:
            raise ValueError("Base: .docx de hasta 20 MB")
        try:
            Document(io.BytesIO(datos))
        except Exception:
            raise ValueError("No es un .docx válido")
        fn = f"{eid}_base.docx"
    if e[clave] and e[clave] != fn:
        (_est_dir() / e[clave]).unlink(missing_ok=True)
    (_est_dir() / fn).write_bytes(datos)
    e[clave] = fn
    (_est_dir() / f"{eid}.json").write_text(json.dumps(e, ensure_ascii=False, indent=1))
    return e


def estilo_borrar(eid):
    e = estilo_get(eid)
    for k in ("logo", "base_docx"):
        if e[k]:
            (_est_dir() / e[k]).unlink(missing_ok=True)
    (_est_dir() / f"{eid}.json").unlink()
    f = P.WORK / "carpetas.json"  # y de las carpetas que lo tenían por defecto
    try:
        j = json.loads(f.read_text())
        if isinstance(j, dict) and any(v == eid for v in j.get("estilos", {}).values()):
            j["estilos"] = {k: v for k, v in j["estilos"].items() if v != eid}
            f.write_text(json.dumps(j, ensure_ascii=False))
    except (OSError, ValueError):
        pass


def estilo_para(c, eid=None):
    """El pedido, o el de la carpeta de la reunión (carpetas.json -> estilos); None si no hay."""
    if eid == "":
        return None
    if not eid:
        try:
            j = json.loads((P.WORK / "carpetas.json").read_text())
            eid = j.get("estilos", {}).get(c["carpeta"]) if isinstance(j, dict) else None
        except (OSError, ValueError):
            eid = None
    try:
        return estilo_get(eid) if eid else None
    except KeyError:
        return None


# ---------- DOCX ----------
def docx_bytes(c, est=None):
    from PIL import Image
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, RGBColor
    est = est or {}
    base = _est_dir() / est["base_docx"] if est.get("base_docx") else None
    doc = Document(str(base)) if base and base.is_file() else Document()
    rgb = RGBColor.from_string(est["color"][1:].upper()) if est.get("color") else None
    D = c["data"]
    ctx = {"{fecha}": f"{c['fecha']:%d/%m/%Y}", "{título}": D["titulo"], "{titulo}": D["titulo"]}

    def sub(txt):
        for k, v in ctx.items():
            txt = txt.replace(k, v)
        return txt

    def con_pagina(par, txt):
        for i, parte in enumerate(sub(txt).split("{página}")):
            if i:
                f = OxmlElement("w:fldSimple")
                f.set(qn("w:instr"), "PAGE")
                r = OxmlElement("w:r")
                t = OxmlElement("w:t")
                t.text = "1"
                r.append(t)
                f.append(r)
                par._p.append(f)
            if parte:
                par.add_run(parte)

    def estilo(p, nombre):
        try:
            p.style = nombre
        except KeyError:
            pass
        return p

    def H(txt, n):
        try:
            h = doc.add_heading(txt, n)
        except KeyError:  # la base no trae estilos de título
            h = doc.add_paragraph()
            h.add_run(txt).bold = True
        if rgb:
            for r in h.runs:
                r.font.color.rgb = rgb
        return h

    def li(txt):
        p = estilo(doc.add_paragraph(), "List Bullet")
        if p.style.name != "List Bullet":
            txt = "• " + txt
        p.add_run(txt)
        return p
    sec = doc.sections[0]
    if est.get("cabecera") or est.get("logo"):
        p = sec.header.paragraphs[0]
        lg = _est_dir() / est["logo"] if est.get("logo") else None
        if lg and lg.is_file():
            p.add_run().add_picture(str(lg), height=Inches(0.5))
            p.add_run("  ")
        con_pagina(p, est.get("cabecera", ""))
    if est.get("pie"):
        con_pagina(sec.footer.paragraphs[0], est["pie"])
    H(f"{c['doc'].capitalize()}: {D['titulo']}", 0).alignment = 0
    doc.add_paragraph(f"Duración {fm(D['duracion'])} · {c['fecha']:%d/%m/%Y}")
    cf = [(n, est["campos"][k]) for k, n in CAMPOS if (est.get("campos") or {}).get(k)]
    if cf:
        t = doc.add_table(rows=0, cols=2)
        estilo(t, "Table Grid")
        for n, v in cf:
            r = t.add_row().cells
            r[0].text, r[1].text = n, sub(v)
            r[0].paragraphs[0].runs[0].bold = True
    for cl, tit, its in bloques(c):
        if cl == "secciones":
            H(tit, 1)
            for s in its:
                H(f"[{fm(s['t'])}] {s['titulo']}", 2)
                doc.add_paragraph(s["resumen"])
                for x in s["puntos"]:
                    li(f"[{fm(x['t'])}] {x['texto']}")
            continue
        H(tit, 1)
        if cl == "texto":
            doc.add_paragraph(its[0]["texto"])
        elif cl == "parrafos":
            for x in its:
                doc.add_paragraph(f"[{fm(x['t'])}] {x['texto']}")
        elif cl == "lista":
            for x in its:
                li(f"[{fm(x['t'])}] {x['texto']}" + (f" ({x['sub']})" if x.get("sub") else ""))
        elif cl == "tareas":
            t = doc.add_table(rows=1, cols=4)
            estilo(t, "Table Grid")
            for cel, n in zip(t.rows[0].cells, ("Tarea", "Responsable", "Fecha / plazo", "Minuto")):
                cel.text = n
                cel.paragraphs[0].runs[0].bold = True
            for x in its:
                f = fecha(x["plazo"], c["fecha"])
                for cel, v in zip(t.add_row().cells, (x["texto"], x["resp"], f"{f:%d/%m/%Y}" if f else x["plazo"], fm(x["t"]))):
                    cel.text = v
        elif cl == "capturas":
            for k in its:
                im = io.BytesIO()  # los JPEG de ffmpeg (marcador COM primero) no los reconoce python-docx: se recodifican
                Image.open(c["d"] / c["img"] / k["img"]).convert("RGB").save(im, "JPEG", quality=88)
                doc.add_picture(im, width=Inches(5.5))
                estilo(doc.add_paragraph(f"[{fm(k['t'])}] {k.get('titulo') or 'Pantalla compartida'}"), "Caption")
        elif cl == "trans":
            for s in its:
                p = doc.add_paragraph(f"[{fm(s['s'])}] ")
                if s.get("spk"):
                    p.add_run(f"{s['spk']}: ").bold = True
                p.add_run(s["text"])
    if est.get("firmas"):
        doc.add_paragraph()
        t = doc.add_table(rows=1, cols=len(est["firmas"]))
        for cel, f in zip(t.rows[0].cells, est["firmas"]):
            cel.text = f"\n\n______________________\n{f['cargo']}\n{f['nombre']}"
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


# ---------- despacho ----------
def exportar(mid, ver, fmt, estilo=None, hablantes=False, evento=False, anon=False, dudas=False):
    """-> (bytes, nombre de fichero, tipo MIME). ValueError: formato desconocido; SinNavegador: pdf sin navegador."""
    if fmt not in FORMATOS:
        raise ValueError(f"Formato desconocido: {fmt}")
    c = carga(mid, ver, anon, dudas)
    est = estilo_para(c, estilo) if fmt == "docx" else None
    if fmt in ("srt", "vtt") and not c["segs"]:
        raise ValueError("Sin transcripción")
    datos = {"docx": lambda: docx_bytes(c, est), "pdf": lambda: pdf_bytes(c), "md": lambda: md(c), "obsidian": lambda: obsidian_zip(c), "notion": lambda: notion_zip(c),
             "ics": lambda: ics(c, evento), "csv_todoist": lambda: csv_todoist(c), "csv_jira": lambda: csv_jira(c),
             "srt": lambda: subtitulos(c["segs"], False, hablantes), "vtt": lambda: subtitulos(c["segs"], True, hablantes)}[fmt]()
    ext, mime = FORMATOS[fmt]
    slug = re.sub(r"[^\w.-]+", "_", c["data"]["titulo"], flags=re.A)[:60].strip("_") or "documento"
    return (datos.encode() if isinstance(datos, str) else datos), f"{slug}.{ext}", mime
