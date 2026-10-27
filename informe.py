"""F6 Informe de estado / comité: KPIs deterministas sobre las entidades del registro, semáforos sugeridos con umbrales fijos (D7, editables
por informe), narrativa con IA opcional (carpetas._redactar: frases con fuentes [n] y control de cifras) y exportación docx|pdf|pptx|md|json|xlsx.

Invariantes:
- Todo es usable sin LLM: kpis + semáforos + narrativa vacía para escribir a mano.
- Una fila de `informe` = una versión. serie = id de la primera versión; «Recalcular» y «Redactar» crean versión nueva (se guardan las
  últimas VERSIONES de cada serie); editar (PATCH) cambia la versión en la que se está, con modificado nuevo.
- semaforos = {umbrales, items: {dimensión: {color, justificacion, sugerido, sugerido_just, manual}}}: manual=True se respeta al recalcular.
- contenido = {kpis (instantánea), narrativa {sección: [{texto, refs:[ref]}]}, notas}. ref = {id, codigo, titulo, revision, cita}.
- Escala de riesgos: palabras alto|medio|bajo -> 4|3|2 y números 1..5 tal cual; severidad = probabilidad×impacto (escala 5×5; rojo >= 12, ámbar >= 6); lo que no se entiende queda «sin valorar» (nunca se infiere).
- Exportar con borrador=False: KPIs recalculados solo con lo aprobado y fuera las frases que citan algo propuesto. Con borrador=True lo no
  revisado lleva la marca BORRADOR (como export_registro).
"""
import datetime as dt
import html
import io
import json
import re
import time

import carpetas
import export_registro as ER
import exportar as X
import registro

VERSIONES = 10
DIMS = ("Global", "Alcance", "Plazo", "Coste", "Riesgos")
COLORES = ("verde", "ambar", "rojo", "gris")
UMBRALES = {"vencidas_ambar": 1, "vencidas_rojo": 3, "sev_ambar": 6, "sev_rojo": 12, "preguntas_ambar": 3, "escalados_rojo": 2, "hitos_vencidos_rojo": 1}
SECCIONES = {"logros": "Logros del periodo", "proximos": "Próximos pasos", "riesgos": "Riesgos principales", "comite": "Decisiones requeridas al comité"}
CERR = {"hecho", "cerrado", "completado", "resuelto", "mitigado"}
FUERA = {"descartado", "cancelado"}
HEX = {"verde": "2E7D32", "ambar": "F9A825", "rojo": "C62828", "gris": "9E9E9E"}
FORMA = {"verde": "●", "ambar": "▲", "rojo": "■", "gris": "○"}  # forma + texto: nunca solo color
NOMBRE = {"verde": "Verde", "ambar": "Ámbar", "rojo": "Rojo", "gris": "Sin datos"}
EXT = {"docx": "docx", "pdf": "pdf", "pptx": "pptx", "md": "md", "json": "json", "xlsx": "xlsx"}
MIME = {**{k: v for k, v in ER.MIME.items() if k in EXT}, "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}


# ---------- fechas y escalas ----------
def _d(s):
    try:
        return dt.date.fromisoformat(str(s)[:10]) if s else None
    except ValueError:
        return None


def _alta(x):
    return _d(x["fecha_alta"]) or dt.date.fromtimestamp(x["creado"] or 0)


def _limite(x):
    if f := _d(x["fecha_limite"]):
        return f
    pl = (x["campos"] or {}).get("plazo")
    return carpetas.plazo_fecha(pl, _alta(x)) if pl else None


def _cierre(x):
    return _d(x["fecha_cierre"]) or dt.date.fromtimestamp(x["modificado"] or x["creado"] or 0)


def _abierta(x, f):
    e = (x["estado"] or "").lower()
    return _alta(x) <= f and e not in FUERA and not (e in CERR and _cierre(x) <= f)


def nivel(v):
    """probabilidad/impacto -> 1..5 o None (sin valorar)."""
    s = str(v or "").strip().lower()
    if s.isdigit():
        return int(s) if 1 <= int(s) <= 5 else None
    m = re.match(r"(muy\s+)?(alt|high|crít|crit|med|baj|low)", s)
    return m and {"alt": 4, "hig": 4, "crí": 4, "cri": 4, "med": 3, "baj": 2, "low": 2}[m[2][:3]]


def _ref(x, docs):
    ci = (x.get("citas") or [None])[0]
    cita = ci and {k: ci[k] for k in ("mid", "ver", "t", "texto")} | {"doc": docs.get(ci["mid"], ("", ""))[0]}
    return {"id": x["id"], "codigo": x["codigo"], "titulo": x["titulo"], "revision": x["revision"], "estado": x["estado"], "dueno": x["dueno"], "cita": cita}


def _bloque(xs, **k):
    return {**k, "propuestos": sum(x["revision"] == "propuesto" for x in xs)}


# ---------- KPIs ----------
def kpis(c, desde, hasta, docs=(), solo_aprobados=False):
    """KPIs deterministas del periodo [desde, hasta] (ISO). Cada bloque lleva `propuestos`: cuántos de sus elementos están sin revisar."""
    desde, hasta = _d(desde), _d(hasta)
    if not desde or not hasta or desde > hasta:
        raise ValueError("Periodo no válido")
    ok = (lambda r: r == "aprobado") if solo_aprobados else (lambda r: r not in registro.OCULTA)
    xs = registro.completas([x for x in registro.listar(c) if ok(x["revision"])])
    nd = {d["mid"]: (d["data"].get("titulo") or d["nombre"], d["fecha"]) for d in docs}
    R = lambda x: _ref(x, nd)
    por = {}
    for x in xs:
        por.setdefault(x["tipo"], []).append(x)
    # tareas
    ts = por.get("tarea", [])
    ab = sorted((x for x in ts if _abierta(x, hasta)), key=lambda x: (_limite(x) is None, _limite(x) or hasta))
    ve = [x for x in ab if _limite(x) and _limite(x) < hasta]
    ce = [x for x in ts if (x["estado"] or "").lower() in CERR and desde <= _cierre(x) <= hasta]
    lim = lambda x: (_limite(x) or "") and _limite(x).isoformat()
    tareas = _bloque(ab + ce, abiertas=len(ab), vencidas=len(ve), cerradas=len(ce), sin_plazo=sum(not _limite(x) for x in ab),
                     lista_abiertas=[{**R(x), "limite": lim(x), "vencida": x in ve} for x in ab[:15]], lista_cerradas=[R(x) for x in ce[:15]])
    # riesgos
    rs = [x for x in por.get("riesgo", []) if (x["estado"] or "").lower() not in CERR | FUERA]
    val = []
    for x in rs:
        p, i = nivel(x["campos"].get("probabilidad")), nivel(x["campos"].get("impacto"))
        val.append({**R(x), "p": p, "i": i, "sev": p * i if p and i else None})
    val.sort(key=lambda r: (r["sev"] is None, -(r["sev"] or 0), r["codigo"] or ""))
    banda = lambda s: "sin_valorar" if s is None else "alta" if s >= UMBRALES["sev_rojo"] else "media" if s >= UMBRALES["sev_ambar"] else "baja"
    sev = {b: sum(banda(r["sev"]) == b for r in val) for b in ("alta", "media", "baja", "sin_valorar")}
    mat = {}
    for r in val:
        if r["sev"]:
            mat.setdefault((r["p"], r["i"]), []).append({"codigo": r["codigo"], "titulo": r["titulo"]})
    riesgos = _bloque(rs, total=len(rs), por_severidad=sev, top=val[:5], matriz=[{"p": p, "i": i, "items": v} for (p, i), v in sorted(mat.items())])
    # decisiones del periodo
    fd = lambda x: _d(x["campos"].get("fecha")) or _alta(x)
    dec = [x for x in por.get("decision", []) if desde <= fd(x) <= hasta]
    decisiones = _bloque(dec, n=len(dec), lista=[{**R(x), "fecha": fd(x).isoformat()} for x in dec])
    est = lambda t: {"total": len(por.get(t, [])), "por_estado": dict(sorted(_cuenta((x["estado"] or "sin estado") for x in por.get(t, [])).items()))}
    # preguntas abiertas y escalados
    pq = [x for x in por.get("pregunta", []) if _abierta(x, hasta)]
    es = [x for x in xs if (x["estado"] or "").lower() == "escalado" or x["campos"].get("escalado")]
    preguntas = _bloque(pq + es, abiertas=len(pq), escalados=len(es), lista=[R(x) for x in pq[:15]], lista_escalados=[R(x) for x in es[:15]])
    hs = sorted(por.get("hito", []), key=lambda x: (_limite(x) is None, _limite(x) or hasta))
    hl = [{**R(x), "fecha": lim(x), "vencido": bool(_abierta(x, hasta) and _limite(x) and _limite(x) < hasta)} for x in hs[:20]]
    hitos = _bloque(hs, lista=hl, vencidos=sum(h["vencido"] for h in hl))
    # serie de tareas abiertas por reunión (hasta el final del periodo; las 12 últimas)
    reus = {d["mid"]: _d(d["fecha"]) for d in docs}
    for x in xs:
        if x["mid"] and x["mid"] not in reus:
            reus[x["mid"]] = _alta(x)
    serie = [{"fecha": f.isoformat(), "mid": m, "doc": nd.get(m, ("", ""))[0], "abiertas": sum(_abierta(x, f) for x in ts)}
             for m, f in sorted(reus.items(), key=lambda kv: kv[1]) if f and f <= hasta][-12:]
    rev = {"propuesto": 0, "aprobado": 0, "por_tipo": {}}
    for x in xs:
        if x["revision"] in ("propuesto", "aprobado"):
            rev[x["revision"]] += 1
            rev["por_tipo"].setdefault(x["tipo"], {"propuesto": 0, "aprobado": 0})[x["revision"]] += 1
    return {"periodo": {"desde": desde.isoformat(), "hasta": hasta.isoformat()}, "solo_aprobados": solo_aprobados, "tareas": tareas, "riesgos": riesgos,
            "decisiones": decisiones, "requisitos": _bloque(por.get("requisito", []), **est("requisito")), "historias": _bloque(por.get("historia", []), **est("historia")),
            "preguntas": preguntas, "hitos": hitos, "serie": serie, "revision": rev}


def _cuenta(it):
    o = {}
    for k in it:
        o[k] = o.get(k, 0) + 1
    return o


# ---------- semáforos ----------
def semaforos_sugeridos(k, umbrales=None):
    """-> {dimensión: {color, justificacion}}. Reglas fijas (D7) con umbrales editables por informe; gris = sin datos."""
    u = {**UMBRALES, **(umbrales or {})}
    t, r, p, h = k["tareas"], k["riesgos"], k["preguntas"], k["hitos"]
    s = {}
    if not t["abiertas"] + t["cerradas"] and not h["lista"]:
        s["Plazo"] = ("gris", "Sin tareas ni hitos en el registro.")
    else:
        v = t["vencidas"]
        txt = f"{v} tareas vencidas de {t['abiertas']} abiertas" + (f"; {h['vencidos']} hitos vencidos" if h["vencidos"] else "") + "."
        s["Plazo"] = ("rojo" if v >= u["vencidas_rojo"] or h["vencidos"] >= u["hitos_vencidos_rojo"] else "ambar" if v >= u["vencidas_ambar"] else "verde", txt)
    if not r["total"]:
        s["Riesgos"] = ("gris", "Sin riesgos abiertos en el registro.")
    else:
        mx = max((x["sev"] or 0 for x in r["top"]), default=0)
        sv = r["por_severidad"]["sin_valorar"]
        txt = f"{r['total']} riesgos abiertos; severidad máxima {mx or 'sin valorar'}" + (f"; {sv} sin valorar" if sv else "") + "."
        s["Riesgos"] = ("rojo" if mx >= u["sev_rojo"] else "ambar" if mx >= u["sev_ambar"] or sv else "verde", txt)
    if not k["requisitos"]["total"] + k["historias"]["total"] + p["abiertas"] + p["escalados"]:
        s["Alcance"] = ("gris", "Sin requisitos, historias ni preguntas abiertas.")
    else:
        txt = f"{k['requisitos']['total']} requisitos, {k['historias']['total']} historias; {p['abiertas']} preguntas abiertas, {p['escalados']} escalados."
        s["Alcance"] = ("rojo" if p["escalados"] >= u["escalados_rojo"] else "ambar" if p["abiertas"] >= u["preguntas_ambar"] or p["escalados"] else "verde", txt)
    s["Coste"] = ("gris", "Sin fuente de datos de coste: valóralo a mano.")
    peor = max((s[d][0] for d in s if s[d][0] != "gris"), key=("verde", "ambar", "rojo").index, default="gris")
    s["Global"] = (peor, "Peor color de las dimensiones con datos." if peor != "gris" else "Sin datos suficientes.")
    return {d: {"color": s[d][0], "justificacion": s[d][1]} for d in DIMS}


def _fundir(prev, k, umbrales):
    sug = semaforos_sugeridos(k, umbrales)
    out = {}
    for d in DIMS:
        p = (prev or {}).get(d) or {}
        man = bool(p.get("manual"))
        out[d] = {"color": p["color"] if man else sug[d]["color"], "justificacion": p.get("justificacion", "") if man else sug[d]["justificacion"],
                  "sugerido": sug[d]["color"], "sugerido_just": sug[d]["justificacion"], "manual": man}
    return out


# ---------- narrativa con LLM ----------
SYS_INF = carpetas.SYS_RED + """
Reglas del informe: escribe directamente los hechos; nada de frases sobre la propia sección o el informe («La sección recoge…», «Este apartado…»). Nunca presentes fechas del registro o de creación como el momento en que ocurrió algo: usa solo fechas que aparezcan en el texto de la fuente (plazos, citas). No infieras nada que la fuente no diga (por ejemplo, no afirmes «sin responsable» si la fuente no lo dice)."""


def _txt(r, pref):
    ex = [f"responsable {r['dueno']}" if r.get("dueno") else "", f"plazo {r['limite']}" if r.get("limite") else "",  # sin la fecha de alta: es del registro, no un hecho dicho
          f"severidad {r['sev']}" if r.get("sev") else "", "vencida" if r.get("vencida") or r.get("vencido") else ""]
    return pref + r["titulo"] + (" (" + ", ".join(e for e in ex if e) + ")" if any(ex) else "")


def _fuente(r, pref):
    return _txt(r, pref), {k: r[k] for k in ("id", "codigo", "titulo", "revision", "cita")}


def fuentes(k):
    """sección -> [(texto, ref)] con lo aprobado y lo propuesto del periodo."""
    t, p = k["tareas"], k["preguntas"]
    return {"logros": [_fuente(r, "Tarea cerrada: ") for r in t["lista_cerradas"]] + [_fuente(r, "Decisión tomada: ") for r in k["decisiones"]["lista"]],
            "proximos": [_fuente(r, "Tarea abierta: ") for r in t["lista_abiertas"]] + [_fuente(r, "Hito: ") for r in k["hitos"]["lista"] if not r["vencido"]],
            "riesgos": [_fuente(r, "Riesgo: ") for r in k["riesgos"]["top"]] + [_fuente(r, "Tarea vencida: ") for r in t["lista_abiertas"] if r["vencida"]],
            "comite": [_fuente(r, "Pregunta abierta: ") for r in p["lista"]] + [_fuente(r, "Escalado: ") for r in p["lista_escalados"]] +
                      [_fuente(r, "Riesgo alto: ") for r in k["riesgos"]["top"] if (r["sev"] or 0) >= UMBRALES["sev_rojo"]]}


def redactar(c, chat, k):
    """-> narrativa {sección: [{texto, refs}]}: una llamada por sección con fuentes; frases sin fuente o con cifras inventadas fuera
    (carpetas._redactar). Sección sin fuentes: vacía (sin llamada)."""
    out = {}
    for s, fu in fuentes(k).items():
        out[s] = carpetas._redactar(chat, fu[:20], f"Informe de estado del proyecto «{c}» ({k['periodo']['desde']} a {k['periodo']['hasta']}). "
                                    f"Sección: {SECCIONES[s]}. Cuenta solo lo de esta sección.", SYS_INF) if fu else []
    return out


# ---------- persistencia ----------
def _fila(r):
    x = dict(r)
    x["semaforos"], x["contenido"] = json.loads(x["semaforos"] or "{}"), json.loads(x["contenido"] or "{}")
    return x


def _pid(db, c):
    p = registro.proyecto(c, False, db)
    if not p:
        raise KeyError("Ese informe no existe")
    return p["id"]


def _get(db, c, iid):
    r = db.execute("SELECT * FROM informe WHERE id=? AND proyecto_id=?", (int(iid), _pid(db, c))).fetchone()
    if not r:
        raise KeyError("Ese informe no existe")
    return _fila(r)


def resumen(x):
    return {k: x[k] for k in ("id", "serie", "version", "titulo", "desde", "hasta", "estado", "creado", "modificado")}


def listar(c):
    """Última versión de cada serie, la más reciente primero."""
    with registro._db() as db:
        p = registro.proyecto(c, False, db)
        if not p:
            return []
        rs = db.execute("SELECT * FROM informe i WHERE proyecto_id=? AND version=(SELECT MAX(version) FROM informe WHERE serie=i.serie) ORDER BY modificado DESC", (p["id"],))
        return [resumen(_fila(r)) for r in rs]


def obtener(c, iid):
    with registro._db() as db:
        x = _get(db, c, iid)
        x["versiones"] = [dict(r) for r in db.execute("SELECT id, version, estado, modificado FROM informe WHERE serie=? ORDER BY version DESC", (x["serie"],))]
        return x


def eliminar(c, iid):
    """Quita esa versión (si era la única de su serie, desaparece la serie). -> id de la última versión que queda en la serie o None."""
    with registro._db(escribe=True) as db:
        x = _get(db, c, iid)
        db.execute("DELETE FROM informe WHERE id=?", (x["id"],))
        r = db.execute("SELECT id FROM informe WHERE serie=? ORDER BY version DESC LIMIT 1", (x["serie"],)).fetchone()
    return r and r[0]


def periodo_defecto(c, hoy=None):
    hoy = hoy or dt.date.today()
    xs = [x for x in listar(c) if _d(x["hasta"])]
    ult = max((_d(x["hasta"]) for x in xs), default=None)
    desde = ult + dt.timedelta(1) if ult and ult < hoy else hoy - dt.timedelta(30)
    return desde.isoformat(), hoy.isoformat()


def crear(c, desde=None, hasta=None, titulo=None, docs=(), hoy=None):
    """Borrador nuevo del periodo (por defecto: desde el último informe o los últimos 30 días) con KPIs y semáforos sugeridos."""
    if not desde or not hasta:
        d0, h0 = periodo_defecto(c, hoy)
        desde, hasta = desde or d0, hasta or h0
    k = kpis(c, desde, hasta, docs)
    sem = {"umbrales": dict(UMBRALES), "items": _fundir(None, k, UMBRALES)}
    con = {"kpis": k, "narrativa": {s: [] for s in SECCIONES}, "notas": ""}
    t = time.time()
    with registro._db(escribe=True) as db:
        p = registro.proyecto(c, db=db)["id"]
        iid = db.execute("INSERT INTO informe(proyecto_id, titulo, desde, hasta, semaforos, contenido, estado, creado, modificado, version) VALUES (?,?,?,?,?,?,?,?,?,1)",
                         (p, titulo or f"Informe de estado {_d(desde):%d/%m/%Y} – {_d(hasta):%d/%m/%Y}", desde, hasta, json.dumps(sem, ensure_ascii=False),
                          json.dumps(con, ensure_ascii=False), "borrador", t, t)).lastrowid
        db.execute("UPDATE informe SET serie=? WHERE id=?", (iid, iid))
    return obtener(c, iid)


def _escribir(db, x, nueva):
    vals = (x["titulo"], x["desde"], x["hasta"], json.dumps(x["semaforos"], ensure_ascii=False), json.dumps(x["contenido"], ensure_ascii=False), x["estado"], time.time())
    if not nueva:
        db.execute("UPDATE informe SET titulo=?, desde=?, hasta=?, semaforos=?, contenido=?, estado=?, modificado=? WHERE id=?", (*vals, x["id"]))
        return x["id"]
    v = db.execute("SELECT MAX(version) FROM informe WHERE serie=?", (x["serie"],)).fetchone()[0] + 1
    iid = db.execute("INSERT INTO informe(proyecto_id, titulo, desde, hasta, semaforos, contenido, estado, modificado, creado, serie, version) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                     (x["proyecto_id"], *vals, x["creado"], x["serie"], v)).lastrowid
    db.execute("DELETE FROM informe WHERE serie=? AND version<=?", (x["serie"], v - VERSIONES))
    return iid


def editar(c, iid, cm):
    """cm: titulo, desde, hasta, estado (borrador|final), notas, narrativa {sección: [{texto, refs}]}, semaforos {dim: {color?, justificacion?,
    manual?}} (manual=False vuelve al sugerido), umbrales {clave: número} (re-sugiere con los KPIs guardados). Misma versión."""
    if not isinstance(cm, dict) or set(cm) - {"titulo", "desde", "hasta", "estado", "notas", "narrativa", "semaforos", "umbrales"}:
        raise ValueError("Cambio no válido")
    with registro._db(escribe=True) as db:
        x = _get(db, c, iid)
        for f in ("titulo", "notas"):
            if f in cm:
                (x if f == "titulo" else x["contenido"])[f] = str(cm[f] or "")
        if "estado" in cm:
            if cm["estado"] not in ("borrador", "final"):
                raise ValueError("Estado no válido")
            x["estado"] = cm["estado"]
        for f in ("desde", "hasta"):
            if f in cm:
                if not _d(cm[f]):
                    raise ValueError("Fecha no válida")
                x[f] = _d(cm[f]).isoformat()
        if "narrativa" in cm:
            n = cm["narrativa"]
            if not isinstance(n, dict) or set(n) - set(SECCIONES) or not all(isinstance(v, list) for v in n.values()):
                raise ValueError("Narrativa no válida")
            ents = {e["id"]: e for e in registro.listar(c, db=db)}

            def refs(f):
                rs = f.get("refs") if isinstance(f, dict) else []
                if rs is None:
                    rs = []
                if not isinstance(rs, list) or not all(isinstance(r, dict) and isinstance(r.get("id"), str) for r in rs):
                    raise ValueError("refs debe ser una lista de {id, …}")
                if falta := next((r["id"] for r in rs if r["id"] not in ents), None):
                    raise KeyError(f"La referencia {falta} no existe en el proyecto")
                # codigo y título salen del registro y la revisión no se guarda: se consulta al exportar (preparar), nunca se fía de lo que manda el cliente
                return [{"id": r["id"], "codigo": ents[r["id"]]["codigo"], "titulo": ents[r["id"]]["titulo"], "cita": r["cita"] if isinstance(r.get("cita"), dict) else {}} for r in rs]
            x["contenido"]["narrativa"] = {**x["contenido"].get("narrativa", {}), **{s: [{"texto": str(f.get("texto") or "") if isinstance(f, dict) else str(f),
                                           "refs": refs(f)} for f in v] for s, v in n.items()}}
        sem = x["semaforos"]
        if "umbrales" in cm:
            if not isinstance(cm["umbrales"], dict) or set(cm["umbrales"]) - set(UMBRALES) or not all(isinstance(v, (int, float)) for v in cm["umbrales"].values()):
                raise ValueError("Umbrales no válidos")
            sem["umbrales"] = {**sem.get("umbrales", UMBRALES), **cm["umbrales"]}
        for d, v in (cm.get("semaforos") or {}).items():
            if d not in DIMS or not isinstance(v, dict) or v.get("color", "verde") not in COLORES:
                raise ValueError("Semáforo no válido")
            if "justificacion" in v and not isinstance(v["justificacion"], str):
                raise ValueError("La justificación debe ser texto")
            it = sem["items"][d]
            if v.get("manual") is False:
                it["manual"] = False
            else:
                it.update({k: v[k] for k in ("color", "justificacion") if k in v}, manual=True)
        sem["items"] = _fundir(sem["items"], x["contenido"]["kpis"], sem.get("umbrales"))
        _escribir(db, x, False)
    return obtener(c, iid)


def recalcular(c, iid, docs=(), narrativa=None):
    """Versión nueva con KPIs del periodo actuales y semáforos re-sugeridos (los manuales se quedan); narrativa: la sustituye."""
    x = obtener(c, iid)
    k = kpis(c, x["desde"], x["hasta"], docs)
    x["contenido"]["kpis"] = k
    if narrativa is not None:
        x["contenido"]["narrativa"] = narrativa
    x["semaforos"]["items"] = _fundir(x["semaforos"]["items"], k, x["semaforos"].get("umbrales"))
    with registro._db(escribe=True) as db:
        iid = _escribir(db, x, True)
    return obtener(c, iid)


def redactar_informe(c, iid, chat, docs=()):
    """Trabajo «informe» de la cola: KPIs al día + narrativa de la IA en una versión nueva."""
    x = obtener(c, iid)
    return recalcular(c, iid, docs, redactar(c, chat, kpis(c, x["desde"], x["hasta"], docs)))


def tras_poblar(c, docs=()):
    """Gancho de poblar.py: el último informe de la carpeta, si es borrador, se pone al día EN LA MISMA versión (KPIs y semáforos sugeridos;
    una pasada automática no crea versiones que empujen fuera las del usuario). Nunca rompe poblar."""
    try:
        xs = listar(c)
        if xs and xs[0]["estado"] == "borrador":
            x = obtener(c, xs[0]["id"])
            x["contenido"]["kpis"] = k = kpis(c, x["desde"], x["hasta"], docs)
            x["semaforos"]["items"] = _fundir(x["semaforos"]["items"], k, x["semaforos"].get("umbrales"))
            with registro._db(escribe=True) as db:
                _escribir(db, x, False)
            return obtener(c, x["id"])
    except Exception as ex:  # ponytail: un fallo del informe no debe tirar la pasada de poblar
        print(f"  informe: no se recalculó el borrador: {ex}", flush=True)
    return None


# ---------- gráficos SVG (los mismos en la app y en el PDF) ----------
PALETA = ["#2a6fdb", "#e08a00", "#2e9d5b", "#b03a8c", "#6b7a8f", "#c0392b", "#17a2b8", "#8e6c00"]


def _e(s):
    return html.escape(str(s), quote=True)


def svg_estados(k):
    """Barras apiladas por estado (requisitos e historias) con leyenda; 100 % del ancho."""
    filas = [(n, k[t]["por_estado"]) for t, n in (("requisitos", "Requisitos"), ("historias", "Historias")) if k[t]["total"]]
    if not filas:
        return ""
    est = list(dict.fromkeys(e for _, d in filas for e in d))
    col = {e: PALETA[i % len(PALETA)] for i, e in enumerate(est)}
    W, x0, o = 600, 90, []
    for j, (n, d) in enumerate(filas):
        y, tot, x = 10 + j * 34, sum(d.values()), x0
        o.append(f'<text x="0" y="{y + 17}" font-size="13" fill="currentColor">{n} ({tot})</text>')
        for e, v in d.items():
            w = (W - x0) * v / tot
            o.append(f'<rect x="{x:.1f}" y="{y}" width="{w:.1f}" height="24" fill="{col[e]}"><title>{_e(n)} · {_e(e)}: {v}</title></rect>')
            if w > 18:
                o.append(f'<text x="{x + w / 2:.1f}" y="{y + 16}" font-size="12" text-anchor="middle" fill="#fff">{v}</text>')
            x += w
    y = 14 + len(filas) * 34
    ley, x = [], 0
    for e in est:
        ley.append(f'<rect x="{x}" y="{y}" width="12" height="12" fill="{col[e]}"/><text x="{x + 16}" y="{y + 11}" font-size="12" fill="currentColor">{_e(e)}</text>')
        x += 24 + 7 * len(e)
    return (f'<svg viewBox="0 0 {W} {y + 20}" role="img" aria-label="Requisitos e historias por estado" class="inf-svg" xmlns="http://www.w3.org/2000/svg">'
            + "".join(o + ley) + "</svg>")


def svg_calor(k):
    """Mapa de calor 5×5 probabilidad × impacto: nº de riesgos por celda y su lista al pasar el ratón (<title>)."""
    m = {(c["p"], c["i"]): c["items"] for c in k["riesgos"]["matriz"]}
    S, x0, o = 46, 40, []
    for p in range(1, 6):
        for i in range(1, 6):
            s, its = p * i, m.get((p, i), [])
            f = "#c62828" if s >= UMBRALES["sev_rojo"] else "#f9a825" if s >= UMBRALES["sev_ambar"] else "#2e7d32"
            x, y = x0 + (i - 1) * S, (5 - p) * S
            tit = f"P{p} × I{i} = {s}: " + ("; ".join(f"{a['codigo']} {a['titulo']}" for a in its) or "sin riesgos")
            o.append(f'<g><title>{_e(tit)}</title><rect x="{x}" y="{y}" width="{S - 2}" height="{S - 2}" fill="{f}" fill-opacity="{0.85 if its else 0.18}"/>'
                     + (f'<text x="{x + S / 2 - 1}" y="{y + S / 2 + 5}" font-size="15" font-weight="700" text-anchor="middle" fill="#fff">{len(its)}</text>' if its else "") + "</g>")
    o += [f'<text x="{x0 + (i - 1) * S + S / 2 - 1}" y="{5 * S + 14}" font-size="11" text-anchor="middle" fill="currentColor">{i}</text>' for i in range(1, 6)]
    o += [f'<text x="{x0 - 8}" y="{(5 - p) * S + S / 2 + 4}" font-size="11" text-anchor="end" fill="currentColor">{p}</text>' for p in range(1, 6)]
    o.append(f'<text x="{x0 + 5 * S / 2}" y="{5 * S + 30}" font-size="12" text-anchor="middle" fill="currentColor">Impacto</text>')
    o.append(f'<text x="12" y="{5 * S / 2}" font-size="12" text-anchor="middle" fill="currentColor" transform="rotate(-90 12 {5 * S / 2})">Probabilidad</text>')
    return f'<svg viewBox="0 0 {x0 + 5 * S + 4} {5 * S + 36}" role="img" aria-label="Mapa de calor de riesgos" class="inf-svg inf-calor" xmlns="http://www.w3.org/2000/svg">{"".join(o)}</svg>'


def svg_linea(k):
    """Tareas abiertas tras cada reunión."""
    s = k["serie"]
    if not s:
        return ""
    W, H, x0, y0 = 600, 180, 34, 150
    mx = max(1, max(p["abiertas"] for p in s))
    X_ = lambda j: x0 + (W - x0 - 20) * (j / max(1, len(s) - 1) if len(s) > 1 else 0.5)
    Y_ = lambda v: y0 - (y0 - 14) * v / mx
    pts = " ".join(f"{X_(j):.1f},{Y_(p['abiertas']):.1f}" for j, p in enumerate(s))
    o = [f'<line x1="{x0}" y1="{y0}" x2="{W - 10}" y2="{y0}" stroke="currentColor" stroke-opacity=".3"/>',
         f'<polyline points="{pts}" fill="none" stroke="#2a6fdb" stroke-width="2.5"/>']
    for j, p in enumerate(s):
        o.append(f'<g><title>{_e(p["fecha"])} {_e(p["doc"])}: {p["abiertas"]} abiertas</title><circle cx="{X_(j):.1f}" cy="{Y_(p["abiertas"]):.1f}" r="4.5" fill="#2a6fdb"/>'
                 f'<text x="{X_(j):.1f}" y="{Y_(p["abiertas"]) - 8:.1f}" font-size="12" text-anchor="middle" fill="currentColor">{p["abiertas"]}</text>'
                 f'<text x="{X_(j):.1f}" y="{y0 + 16}" font-size="11" text-anchor="middle" fill="currentColor">{_d(p["fecha"]):%d/%m}</text></g>')
    return f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="Tareas abiertas por reunión" class="inf-svg" xmlns="http://www.w3.org/2000/svg">{"".join(o)}</svg>'


def graficos(k):
    return {"estados": svg_estados(k), "calor": svg_calor(k), "linea": svg_linea(k)}


# ---------- modelo para exportar ----------
def _numerar(narr):
    """Refs de toda la narrativa -> {id: n} por orden de aparición."""
    n = {}
    for fs in narr.values():
        for f in fs:
            for r in f.get("refs") or []:
                n.setdefault(r["id"], (len(n) + 1, r))
    return n


def preparar(c, x, borrador=True, docs=()):
    """Lo que se exporta: kpis (de la instantánea o, solo aprobados, recalculados), narrativa filtrada, refs numeradas, nº de no revisados."""
    k = x["contenido"]["kpis"] if borrador else kpis(c, x["desde"], x["hasta"], docs, True)
    ents = {e["id"]: e for e in registro.listar(c)}
    fresca = lambda r: {**r, "revision": ents[r["id"]]["revision"], **{q: ents[r["id"]][q] for q in ("codigo", "titulo")}} if r["id"] in ents else {**r, "revision": "descartado"}
    ok = lambda r: r["revision"] in ("propuesto", "aprobado") if borrador else r["revision"] == "aprobado"  # solo aprobados: la revisión sale de la BD, no de lo guardado
    narr = {s: [{**f, "refs": [fresca(r) for r in f.get("refs") or []]} for f in x["contenido"].get("narrativa", {}).get(s, [])] for s in SECCIONES}
    narr = {s: [f for f in fs if all(ok(r) for r in f["refs"])] for s, fs in narr.items()}
    num = _numerar(narr)
    nb = k["revision"]["propuesto"] if borrador else 0
    return {"k": k, "narr": narr, "num": num, "nb": nb, "sem": x["semaforos"]["items"]}


def _frase(f, num):
    return f["texto"] + "".join(f" [{num[r['id']][0]}]" for r in f.get("refs") or [] if r["id"] in num)


def _ref_txt(r):
    ci = r.get("cita") or {}
    return " · ".join(filter(None, [f"{r.get('codigo') or ''} {r.get('titulo') or ''}".strip(), f"{ci.get('doc') or ''} [{X.fm(ci['t'])}]" if ci.get("t") is not None and ci.get("mid") else "",
                                    f"«{ci['texto']}»" if ci.get("texto") else "", ER.BORR if r.get("revision") == "propuesto" else ""]))


def _kpi_filas(k):
    t, r, p = k["tareas"], k["riesgos"], k["preguntas"]
    marca = lambda b: f" (incl. {b['propuestos']} sin revisar)" if b["propuestos"] else ""
    return [["Tareas abiertas", str(t["abiertas"]) + marca(t)], ["Tareas vencidas", str(t["vencidas"])], ["Tareas cerradas en el periodo", str(t["cerradas"])],
            ["Riesgos abiertos", f"{r['total']} (alta {r['por_severidad']['alta']}, media {r['por_severidad']['media']}, baja {r['por_severidad']['baja']}, "
                                 f"sin valorar {r['por_severidad']['sin_valorar']})" + marca(r)],
            ["Decisiones del periodo", str(k["decisiones"]["n"]) + marca(k["decisiones"])], ["Preguntas abiertas", str(p["abiertas"]) + marca(p)], ["Escalados", str(p["escalados"])],
            ["Requisitos", f"{k['requisitos']['total']} · " + ", ".join(f"{e} {n}" for e, n in k["requisitos"]["por_estado"].items()) + marca(k["requisitos"])],
            ["Historias", f"{k['historias']['total']} · " + ", ".join(f"{e} {n}" for e, n in k["historias"]["por_estado"].items()) + marca(k["historias"])],
            ["Hitos", f"{len(k['hitos']['lista'])} ({k['hitos']['vencidos']} vencidos)"], ["Revisión", f"{k['revision']['aprobado']} aprobados · {k['revision']['propuesto']} propuestos"]]


def _top_filas(k):
    v = lambda n: "—" if n is None else str(n)
    return [[r["codigo"] or "", r["titulo"] + (f" [{ER.BORR}]" if r["revision"] == "propuesto" else ""), v(r["p"]), v(r["i"]), v(r["sev"]), r["dueno"] or ""] for r in k["riesgos"]["top"]]


TOP_CAB = ["Código", "Riesgo", "P", "I", "Severidad", "Responsable"]


def _cabecera(c, x, P_):
    return [f"Proyecto: {c or 'Sin carpeta'}", f"Periodo: {_d(x['desde']):%d/%m/%Y} – {_d(x['hasta']):%d/%m/%Y}", f"Estado: {'Final' if x['estado'] == 'final' else 'Borrador'} · versión {x['version']}"], \
        (f"{ER.BORR.upper()}: incluye {P_['nb']} elementos propuestos aún no revisados." if P_["nb"] else "")


def a_md(c, x, P_):
    cab, aviso = _cabecera(c, x, P_)
    o = [f"# {x['titulo']}", "", " · ".join(cab), ""] + ([f"> **{aviso}**", ""] if aviso else [])
    o += ["## Semáforos", "", "| Dimensión | Estado | Justificación |", "|---|---|---|"]
    o += [f"| {d} | {FORMA[s['color']]} {NOMBRE[s['color']]} | {s['justificacion'].replace('|', '/')} |" for d, s in P_["sem"].items()]
    o += ["", "## Indicadores", "", "| Indicador | Valor |", "|---|---|"] + [f"| {a} | {b} |" for a, b in _kpi_filas(P_["k"])]
    o += ["", "## Riesgos principales (top 5)", "", "| " + " | ".join(TOP_CAB) + " |", "|" + "---|" * len(TOP_CAB)] + ["| " + " | ".join(f) + " |" for f in _top_filas(P_["k"])]
    for s, n in SECCIONES.items():
        o += ["", f"## {n}", ""] + ([f"- {_frase(f, P_['num'])}" for f in P_["narr"][s]] or ["—"])
    if x["contenido"].get("notas"):
        o += ["", "## Notas", "", x["contenido"]["notas"]]
    if P_["num"]:
        o += ["", "## Fuentes", ""] + [f"[{n}] {_ref_txt(r)}  " for n, r in P_["num"].values()]
    return "\n".join(o) + "\n"


def a_docx(c, x, P_, est):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    B = X.doc_base(est, {"{fecha}": f"{dt.date.today():%d/%m/%Y}", "{título}": x["titulo"], "{titulo}": x["titulo"]})
    doc = B.doc
    cab, aviso = _cabecera(c, x, P_)
    B.H(x["titulo"], 0)
    doc.add_paragraph(" · ".join(cab))
    if aviso:
        r = doc.add_paragraph().add_run(aviso)
        r.bold, r.font.color.rgb = True, RGBColor(0xB0, 0x20, 0x20)

    def tabla(cabs, filas, sombra=None):
        t = doc.add_table(rows=1, cols=len(cabs))
        B.estilo(t, "Table Grid")
        for cel, n in zip(t.rows[0].cells, cabs):
            cel.text = n
            cel.paragraphs[0].runs[0].bold = True
        for j, f in enumerate(filas):
            cs = t.add_row().cells
            for cel, v in zip(cs, f):
                cel.text = v
            if sombra and (col := sombra(j)):
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), HEX[col])
                cs[1]._tc.get_or_add_tcPr().append(shd)
                for r in cs[1].paragraphs[0].runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0, 0, 0) if col == "ambar" else RGBColor(0xFF, 0xFF, 0xFF)
    sem = list(P_["sem"].items())
    B.H("Semáforos", 1)
    tabla(["Dimensión", "Estado", "Justificación"], [[d, f"{FORMA[s['color']]} {NOMBRE[s['color']]}", s["justificacion"]] for d, s in sem], lambda j: sem[j][1]["color"])
    B.H("Indicadores", 1)
    tabla(["Indicador", "Valor"], _kpi_filas(P_["k"]))
    B.H("Riesgos principales (top 5)", 1)
    tabla(TOP_CAB, _top_filas(P_["k"])) if P_["k"]["riesgos"]["top"] else doc.add_paragraph("Sin riesgos abiertos.")
    for s, n in SECCIONES.items():
        B.H(n, 1)
        for f in P_["narr"][s] or [{"texto": "—"}]:
            B.li(_frase(f, P_["num"]))
    if x["contenido"].get("notas"):
        B.H("Notas", 1)
        doc.add_paragraph(x["contenido"]["notas"])
    if P_["num"]:
        B.H("Fuentes", 1)
        for n, r in P_["num"].values():
            doc.add_paragraph(f"[{n}] {_ref_txt(r)}")
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


CSS = ("@page{size:A4 landscape;margin:14mm}.aviso{background:#ffe58a;color:#000;padding:8px 12px;border-radius:6px;font-weight:600}"
       ".sem td.c{font-weight:700;color:#fff;white-space:nowrap}.sem td.ambar{color:#000}.graf{display:flex;flex-wrap:wrap;gap:18px;align-items:flex-start}"
       ".graf figure{flex:1 1 300px;margin:0;break-inside:avoid}.inf-svg{width:100%;height:auto;color:var(--ink)}.inf-calor{max-width:300px}.borr{color:#b02020;font-weight:700}")


def a_html(c, x, P_, est):
    e = html.escape
    cab, aviso = _cabecera(c, x, P_)
    tab = lambda cabs, fs: "<table><tr>" + "".join(f"<th>{e(h)}</th>" for h in cabs) + "</tr>" + "".join("<tr>" + "".join(f"<td>{e(v)}</td>" for v in f) + "</tr>" for f in fs) + "</table>"
    g = graficos(P_["k"])
    h = [f"<h1>{e(x['titulo'])}</h1><p class=meta>{e(' · '.join(cab))}</p>", f"<p class=aviso>{e(aviso)}</p>" if aviso else "",
         "<h2>Semáforos</h2><table class=sem><tr><th>Dimensión</th><th>Estado</th><th>Justificación</th></tr>" +
         "".join(f'<tr><td>{d}</td><td class="c {s["color"]}" style="background:#{HEX[s["color"]]}">{FORMA[s["color"]]} {NOMBRE[s["color"]]}</td><td>{e(s["justificacion"])}</td></tr>'
                 for d, s in P_["sem"].items()) + "</table>",
         "<h2>Indicadores</h2>" + tab(["Indicador", "Valor"], _kpi_filas(P_["k"])),
         "<div class=graf>" + "".join(f"<figure>{g[k]}<figcaption>{t}</figcaption></figure>" for k, t in (("estados", "Requisitos e historias por estado"),
                                                                                                        ("calor", "Riesgos: probabilidad × impacto"), ("linea", "Tareas abiertas por reunión")) if g[k]) + "</div>",
         "<h2>Riesgos principales (top 5)</h2>" + (tab(TOP_CAB, _top_filas(P_["k"])) if P_["k"]["riesgos"]["top"] else "<p>Sin riesgos abiertos.</p>")]
    for s, n in SECCIONES.items():
        h.append(f"<h2>{n}</h2><ul>" + ("".join(f"<li>{e(_frase(f, P_['num']))}</li>" for f in P_["narr"][s]) or "<li>—</li>") + "</ul>")
    if x["contenido"].get("notas"):
        h.append(f"<h2>Notas</h2><p style='white-space:pre-line'>{e(x['contenido']['notas'])}</p>")
    if P_["num"]:
        h.append("<h2>Fuentes</h2><ol class=meta>" + "".join(f"<li value={n}>{e(_ref_txt(r))}</li>" for n, r in P_["num"].values()) + "</ol>")
    col = f":root{{--acc:{est['color']};--acc-d:{est['color']}}}" if est and est.get("color") else ""
    return (f'<!doctype html><html lang="es"><head><meta charset="utf-8"><title>{e(x["titulo"])}</title><style>{X.CSS}{CSS}{col}body{{max-width:none}}</style></head>'
            f'<body>{X.cabecera_estilo(est)}{"".join(h)}</body></html>')


def a_xlsx(x, P_):
    k = P_["k"]
    R = lambda r: [r.get("codigo") or "", r.get("titulo") or "", r.get("estado") or "", r.get("dueno") or "", ER.BORR if r.get("revision") == "propuesto" else "Aprobado"]
    cab = ["Código", "Título", "Estado", "Responsable", "Revisión"]
    return ER.xlsx({"Semáforos": (["Dimensión", "Color", "Justificación", "Sugerido", "Manual"], [[d, NOMBRE[s["color"]], s["justificacion"], NOMBRE[s["sugerido"]], "sí" if s["manual"] else "no"] for d, s in P_["sem"].items()]),
                    "Indicadores": (["Indicador", "Valor"], _kpi_filas(k)),
                    "Riesgos": (TOP_CAB, _top_filas(k)),
                    "Tareas abiertas": (cab + ["Plazo", "Vencida"], [R(r) + [r["limite"], "sí" if r["vencida"] else ""] for r in k["tareas"]["lista_abiertas"]]),
                    "Decisiones": (cab + ["Fecha"], [R(r) + [r["fecha"]] for r in k["decisiones"]["lista"]]),
                    "Narrativa": (["Sección", "Frase"], [[SECCIONES[s], _frase(f, P_["num"])] for s in SECCIONES for f in P_["narr"][s]]),
                    "Fuentes": (["n", "Fuente"], [[str(n), _ref_txt(r)] for n, r in P_["num"].values()])})


def a_pptx(c, x, P_):
    import pptx_min as PM
    cab, aviso = _cabecera(c, x, P_)
    k = P_["k"]
    lis = lambda s: [_frase(f, P_["num"]) for f in P_["narr"][s]] or ["—"]
    sl = [PM.portada(x["titulo"], cab, aviso),
          PM.semaforos("Semáforos", [(d, s["color"], HEX[s["color"]], f"{FORMA[s['color']]} {NOMBRE[s['color']]}", s["justificacion"]) for d, s in P_["sem"].items()]),
          PM.tabla("Indicadores", ["Indicador", "Valor"], _kpi_filas(k)),
          PM.tabla("Riesgos principales (top 5)", TOP_CAB, _top_filas(k) or [["—"] * len(TOP_CAB)]),
          PM.lista(SECCIONES["comite"], lis("comite")), PM.lista(SECCIONES["proximos"], lis("proximos"))]
    return PM.pptx(sl)


def exportar(c, iid, fmt, borrador=True, docs=(), estilo=None):
    """-> (bytes, nombre, mime). ValueError: formato; exportar.SinNavegador: pdf sin navegador."""
    if fmt not in EXT:
        raise ValueError(f"Formato desconocido: {fmt}")
    x = obtener(c, iid)
    P_ = preparar(c, x, borrador, docs)
    est = X.estilo_para({"carpeta": c}, estilo) if fmt in ("docx", "pdf") else None
    if fmt == "md":
        d = a_md(c, x, P_)
    elif fmt == "json":
        d = json.dumps({"proyecto": c, "informe": resumen(x), "borrador": borrador, "semaforos": P_["sem"], "umbrales": x["semaforos"].get("umbrales"), "kpis": P_["k"],
                        "narrativa": P_["narr"], "notas": x["contenido"].get("notas", ""), "fuentes": {n: r for n, r in P_["num"].values()}}, ensure_ascii=False, indent=1)
    elif fmt == "docx":
        d = a_docx(c, x, P_, est)
    elif fmt == "pdf":
        d = X.pdf_html(a_html(c, x, P_, est))
    elif fmt == "xlsx":
        d = a_xlsx(x, P_)
    else:
        d = a_pptx(c, x, P_)
    slug = re.sub(r'[\\/:*?"<>|\n]+', "_", c).strip() or "proyecto"
    return (d.encode() if isinstance(d, str) else d), f"{slug}-informe-{x['hasta']}.{EXT[fmt]}", MIME[fmt]
