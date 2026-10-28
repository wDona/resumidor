"""F6 informe.py: KPIs con fechas fijas, semáforos, redacción con chat falso, versiones, API y exportaciones. Sin LLM ni red."""
import io
import json
import os
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import informe
import registro
from test_carpetas import C, cl  # noqa: F401

D, H = "2026-03-01", "2026-03-31"


@pytest.fixture()
def reg(cl):
    t, _ = cl
    mk = lambda tipo, titulo, rev="aprobado", **f: registro.crear(C, tipo, titulo=titulo, revision=rev, origen="manual", fecha_alta=f.pop("alta", "2026-03-02"), **f)
    x = {}
    x["t1"] = mk("tarea", "Enviar presupuesto", fecha_limite="2026-03-10", dueno="Ana")                       # vencida
    x["t2"] = mk("tarea", "Revisar contrato", "propuesto", fecha_limite="2026-04-15")                         # abierta, no vencida
    x["t3"] = mk("tarea", "Pintar fachada", estado="hecho", fecha_cierre="2026-03-20")                        # cerrada en el periodo
    x["t4"] = mk("tarea", "Tarea vieja cerrada", estado="hecho", fecha_cierre="2026-01-20", alta="2026-01-02")  # cerrada antes
    x["t5"] = mk("tarea", "Descartada", "descartado", fecha_limite="2026-03-03")
    x["r1"] = mk("riesgo", "Retraso del proveedor", campos={"probabilidad": "alta", "impacto": "alto"}, dueno="Luis")  # 16
    x["r2"] = mk("riesgo", "Lluvias", "propuesto", campos={"probabilidad": "media", "impacto": "baja"})             # 6
    x["r3"] = mk("riesgo", "Sin valorar", campos={"probabilidad": "—"})
    x["d1"] = mk("decision", "Usar teja árabe", alta="2026-03-05")
    x["d2"] = mk("decision", "Decisión de febrero", alta="2026-02-05")
    x["q1"] = mk("pregunta", "¿Quién firma el acta?", "propuesto")
    mk("requisito", "Accesible", estado="abierto"), mk("requisito", "Rampa", estado="hecho"), mk("historia", "Como vecino…", "propuesto", estado="")
    mk("hito", "Fin de obra", fecha_limite="2026-03-15")
    return t, x


def test_kpis(reg):
    x = reg[1]
    k = informe.kpis(C, D, H)
    t = k["tareas"]
    assert (t["abiertas"], t["vencidas"], t["cerradas"], t["propuestos"]) == (2, 1, 1, 1)
    assert t["lista_abiertas"][0]["titulo"] == "Enviar presupuesto" and t["lista_abiertas"][0]["vencida"]
    r = k["riesgos"]
    assert r["total"] == 3 and r["por_severidad"] == {"alta": 1, "media": 1, "baja": 0, "sin_valorar": 1}
    assert [x["sev"] for x in r["top"]] == [16, 6, None] and {"p": 4, "i": 4, "items": [{"id": x["r1"]["id"], "codigo": "R-001", "titulo": "Retraso del proveedor", "sev": 16, "dueno": "Luis"}]} in r["matriz"]
    assert k["decisiones"]["n"] == 1 and k["preguntas"]["abiertas"] == 1 and k["preguntas"]["propuestos"] == 1
    assert k["requisitos"]["por_estado"] == {"abierto": 1, "hecho": 1} and k["historias"]["por_estado"] == {"sin estado": 1}
    assert k["hitos"]["vencidos"] == 1 and k["revision"]["propuesto"] == 4
    se = informe.kpis(C, D, H, docs=[{"mid": "m1", "nombre": "R1", "fecha": "2026-03-09", "data": {}}, {"mid": "m2", "nombre": "R2", "fecha": "2026-03-25", "data": {}}])["serie"]
    assert [(p["doc"], p["abiertas"]) for p in se] == [("R1", 3), ("R2", 2)] and informe.svg_linea({"serie": se}).count("<circle") == 2
    ap = informe.kpis(C, D, H, solo_aprobados=True)
    assert ap["tareas"]["abiertas"] == 1 and ap["riesgos"]["total"] == 2 and ap["revision"]["propuesto"] == 0
    with pytest.raises(ValueError):
        informe.kpis(C, H, D)
    assert informe.nivel("Muy alta") == 4 and informe.nivel("media") == 3 and informe.nivel("baja") == 2 and informe.nivel("4") == 4 and informe.nivel("9") is None and informe.nivel("quizá") is None


def test_semaforos(reg):
    k = informe.kpis(C, D, H)
    s = informe.semaforos_sugeridos(k)
    assert s["Riesgos"]["color"] == "rojo" and s["Plazo"]["color"] == "rojo"  # hito vencido
    assert s["Coste"]["color"] == "gris" and s["Global"]["color"] == "rojo" and "16" in s["Riesgos"]["justificacion"]
    s2 = informe.semaforos_sugeridos(k, {"sev_rojo": 20, "hitos_vencidos_rojo": 5})
    assert s2["Riesgos"]["color"] == "ambar" and s2["Plazo"]["color"] == "ambar"
    vacio = informe.kpis("Nadie", D, H)
    assert all(v["color"] == "gris" for v in informe.semaforos_sugeridos(vacio).values())


def chat_falso(llamadas):
    def chat(sis, user, esq, n, temperature=0):
        llamadas.append(user)
        if "Logros" in user:
            return {"frases": ["Se pintó la fachada [1].", "Se ahorraron 5000 euros [1].", "Frase sin fuente.", "Se decidió la teja árabe [2][99]."]}
        return {"frases": ["Hay que cerrar lo pendiente [1]."]}
    return chat


def test_redactar(reg):
    k = informe.kpis(C, D, H)
    ll = []
    n = informe.redactar(C, chat_falso(ll), k)
    assert len(ll) == 4 and set(n) == set(informe.SECCIONES)
    lg = n["logros"]
    assert [f["texto"] for f in lg] == ["Se pintó la fachada.", "Se decidió la teja árabe."]  # cifra inventada y frase sin fuente fuera
    assert lg[0]["refs"][0]["codigo"] == "T-003" and lg[1]["refs"][0]["titulo"] == "Usar teja árabe" and len(lg[1]["refs"]) == 1


def test_versiones_y_edicion(reg):
    t, x = reg
    i = informe.crear(C, D, H)
    assert i["version"] == 1 and i["serie"] == i["id"] and i["semaforos"]["items"]["Riesgos"]["color"] == "rojo"
    e = informe.editar(C, i["id"], {"semaforos": {"Coste": {"color": "verde", "justificacion": "En presupuesto"}}, "notas": "ojo", "estado": "final",
                                    "narrativa": {"logros": [{"texto": "A mano", "refs": []}]}})
    assert e["semaforos"]["items"]["Coste"] == {"color": "verde", "justificacion": "En presupuesto", "sugerido": "gris", "sugerido_just": e["semaforos"]["items"]["Coste"]["sugerido_just"], "manual": True}
    assert e["contenido"]["notas"] == "ojo" and e["version"] == 1
    registro.editar(x["r1"]["id"], {"estado": "cerrado"})
    r = informe.recalcular(C, i["id"])
    assert r["version"] == 2 and r["serie"] == i["id"] and r["semaforos"]["items"]["Coste"]["color"] == "verde"
    assert r["semaforos"]["items"]["Riesgos"]["color"] == "ambar" and r["contenido"]["narrativa"]["logros"][0]["texto"] == "A mano"
    assert [v["version"] for v in r["versiones"]] == [2, 1] and [x["id"] for x in informe.listar(C)] == [r["id"]]
    for _ in range(12):
        r = informe.recalcular(C, r["id"])
    assert len(r["versiones"]) == informe.VERSIONES and r["versiones"][-1]["version"] == 5
    e = informe.editar(C, r["id"], {"umbrales": {"sev_ambar": 1}})
    assert e["semaforos"]["umbrales"]["sev_ambar"] == 1
    with pytest.raises(ValueError):
        informe.editar(C, r["id"], {"estado": "x"})
    with pytest.raises(ValueError):
        informe.editar(C, r["id"], {"semaforos": {"Coste": {"color": "morado"}}})
    with pytest.raises(KeyError):
        informe.obtener("Otra", r["id"])
    assert informe.periodo_defecto(C)[0] == "2026-04-01" or informe.periodo_defecto(C)[0] > H


def test_api(reg):
    t, _ = reg
    r = t.post("/api/reg/informes", json={"c": C, "desde": D, "hasta": H})
    assert r.status_code == 200
    j = r.json()
    assert j["graficos"]["calor"].startswith("<svg") and "Retraso del proveedor" in j["graficos"]["calor"] and j["proyecto"]["carpeta"] == C
    i = j["id"]
    assert t.get("/api/reg/informes", params={"c": C}).json()["items"][0]["id"] == i
    assert t.get(f"/api/reg/informes/{i}", params={"c": C}).json()["titulo"].startswith("Informe de estado")
    p = t.patch(f"/api/reg/informes/{i}", json={"c": C, "cambios": {"titulo": "Comité marzo"}})
    assert p.status_code == 200 and p.json()["titulo"] == "Comité marzo"
    assert t.patch(f"/api/reg/informes/{i}", json={"c": C, "cambios": {"nada": 1}}).status_code == 400
    r2 = t.post(f"/api/reg/informes/{i}/recalcular", json={"c": C}).json()
    assert r2["version"] == 2
    assert t.get("/api/reg/informes/999", params={"c": C}).status_code == 404
    assert t.get(f"/api/reg/informes/{i}", params={"c": "Nope"}).status_code == 404
    assert t.post("/api/reg/informes", json={"c": C, "desde": H, "hasta": D}).status_code == 400
    d = t.post("/api/reg/informes", json={"c": C}).json()
    assert d["desde"] > H or d["desde"] >= "2026"


def test_exportar(reg):
    t, x = reg
    i = t.post("/api/reg/informes", json={"c": C, "desde": D, "hasta": H}).json()["id"]
    nar = {"logros": [{"texto": "Se pintó", "refs": [{"id": x["t3"]["id"], "codigo": "T-003", "titulo": "Pintar fachada", "revision": "aprobado", "cita": None}]}],
           "comite": [{"texto": "Decidir firma", "refs": [{"id": x["q1"]["id"], "codigo": "P-001", "titulo": "¿Quién firma el acta?", "revision": "propuesto", "cita": None}]}]}
    t.patch(f"/api/reg/informes/{i}", json={"c": C, "cambios": {"narrativa": nar}})
    g = lambda fmt, b=1: t.get(f"/api/reg/informes/{i}/exportar", params={"c": C, "fmt": fmt, "borrador": b})
    m = g("md").text
    assert "BORRADOR" in m and "Se pintó [1]" in m and "Decidir firma [2]" in m and "■ Rojo" in m
    m0 = g("md", 0).text
    assert "BORRADOR" not in m0 and "Decidir firma" not in m0 and "Se pintó [1]" in m0
    j = g("json").json()
    assert j["kpis"]["tareas"]["vencidas"] == 1 and j["borrador"] is True
    doc = Document(io.BytesIO(g("docx").content))
    tx = "\n".join(p.text for p in doc.paragraphs)
    assert "Semáforos" in tx and "BORRADOR" in tx and doc.tables[0].rows[1].cells[0].text == "Global"
    assert 'w:fill="C62828"' in doc.tables[0]._tbl.xml
    z = zipfile.ZipFile(io.BytesIO(g("pptx").content))
    for n in z.namelist():
        if n.endswith((".xml", ".rels")):
            ET.fromstring(z.read(n))
    sl = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
    assert len(sl) == 6 and "Retraso del proveedor" in z.read("ppt/slides/slide4.xml").decode() and "Decidir firma" in z.read("ppt/slides/slide5.xml").decode()
    try:
        import pptx
        assert len(pptx.Presentation(io.BytesIO(g("pptx").content)).slides) == 6
    except ImportError:
        pass
    xz = zipfile.ZipFile(io.BytesIO(g("xlsx").content))
    assert "Retraso del proveedor" in xz.read("xl/worksheets/sheet3.xml").decode()
    assert g("xx").status_code == 400
    if os.environ.get("ACTAS_TEST_PDF") == "1":
        assert g("pdf").content[:4] == b"%PDF"
    html = informe.a_html(C, informe.obtener(C, i), informe.preparar(C, informe.obtener(C, i)), None)
    assert html.count("<svg") == 2 and "landscape" in html


def test_gancho_poblar(reg):
    import poblar
    assert poblar.borrador_informe(C, {}) is None
    i = informe.crear(C, D, H)
    x = poblar.borrador_informe(C, {})
    assert (x["id"], x["version"], len(x["versiones"])) == (i["id"], 1, 1)  # al día en la misma versión: las pasadas no crean versiones
    informe.editar(C, informe.listar(C)[0]["id"], {"estado": "final"})
    assert poblar.borrador_informe(C, {}) is None


def test_redactar_reglas_en_el_prompt(reg):
    vistos = []
    informe.redactar(C, lambda sis, *a, **k: vistos.append(sis) or {"frases": ["Hecho [1]."]}, informe.kpis(C, D, H))
    assert vistos and all("La sección recoge" in s and "fechas del registro" in s and "No infieras" in s for s in vistos)
    assert all("fecha 2026" not in u for s_, u in [x for sec in informe.fuentes(informe.kpis(C, D, H)).values() for x in sec])


def test_eliminar_version(reg):
    t, _ = reg
    i = t.post("/api/reg/informes", json={"c": C, "desde": D, "hasta": H}).json()["id"]
    v2 = t.post(f"/api/reg/informes/{i}/recalcular", json={"c": C}).json()["id"]
    assert t.delete(f"/api/reg/informes/{v2}", params={"c": C}).json() == {"ok": True, "id": i}
    assert t.get(f"/api/reg/informes/{v2}", params={"c": C}).status_code == 404 and [x["id"] for x in informe.listar(C)] == [i]
    assert t.delete(f"/api/reg/informes/{i}", params={"c": C}).json() == {"ok": True, "id": None}
    assert informe.listar(C) == [] and t.delete(f"/api/reg/informes/{i}", params={"c": C}).status_code == 404
