"""Pruebas de /api/reg/exportar (export_registro.py) con TestClient sobre la carpeta sintética de test_carpetas. Sin LLM ni red."""
import csv
import io
import os
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document

import exportar
import registro
from test_carpetas import C, cl  # noqa: F401

CRIT = [{"escenario": "Reserva válida", "dado": "un socio con sesión", "cuando": "reserva una pista", "entonces": "recibe la confirmación"},
        {"escenario": "Pista ocupada", "dado": "una pista reservada", "cuando": "otro la reserva", "entonces": "ve un aviso", "propuesta": True}]


@pytest.fixture()
def reg(cl):
    t, _ = cl
    mk = lambda tipo, titulo, rev, **f: registro.crear(C, tipo, titulo=titulo, revision=rev, origen="manual", **f)
    r1 = mk("requisito", "Reservar pistas", "aprobado", detalle="El socio reserva una pista", prioridad="must", campos={"tipo_req": "funcional", "solicitante": "Laura"})
    mk("requisito", "Pagar online", "propuesto", prioridad="should")
    mk("requisito", "Requisito descartado", "descartado")
    h1 = mk("historia", "Reservar desde el móvil", "propuesto", campos={"como": "socio", "quiero": "reservar", "para": "jugar", "criterios": CRIT, "puntos": 3, "epica": "Reservas"})
    mk("historia", "Cancelar reserva", "aprobado", campos={"criterios": CRIT[:1]})
    mk("historia", "Historia descartada", "descartado")
    registro.relacionar(C, h1["id"], r1["id"], "deriva_de")
    mk("interesado", "Marta Gómez", "aprobado", campos={"rol": "presidenta", "poder": 5, "interes": 5}, detalle="Sin notas")
    mk("interesado", "Pedro Ruiz", "propuesto", campos={"poder": 1, "interes": 2})
    mk("interesado", "Borrado Nombre", "descartado")
    d1 = mk("decision", "Usar Postgres", "aprobado", detalle="Elegir BD", campos={"decisor": "Ana", "alternativas": "MySQL"})
    d2 = mk("decision", "Usar SQLite", "propuesto")
    registro.relacionar(C, d2["id"], d1["id"], "sustituye")
    return t


def g(t, vista, fmt, borrador=1):
    return t.get("/api/reg/exportar", params={"c": C, "vista": vista, "fmt": fmt, "borrador": borrador})


def rows(r):
    return list(csv.DictReader(io.StringIO(r.text.lstrip("﻿"))))


def test_errores(reg):
    assert g(reg, "nada", "csv").status_code == 400
    assert g(reg, "requisitos", "xx").status_code == 400
    assert g(reg, "interesados", "feature").status_code == 400 and g(reg, "decisiones", "csv_jira").status_code == 400
    assert reg.get("/api/reg/exportar", params={"c": "Nope", "vista": "requisitos", "fmt": "csv"}).status_code == 404


def test_csv_json_md_borrador(reg):
    r = g(reg, "requisitos", "csv")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv") and "requisitos.csv" in r.headers["content-disposition"]
    assert r.text.startswith("﻿") and "\r\n" in r.text
    filas = rows(r)
    assert [f["Código"] for f in filas] == ["REQ-001", "REQ-002"] and filas[1]["Revisión"].startswith("BORRADOR") and filas[0]["Revisión"] == "Aprobado"
    assert "descartado" not in r.text
    ap = g(reg, "requisitos", "csv", 0)
    assert [f["Código"] for f in rows(ap)] == ["REQ-001"] and "BORRADOR" not in ap.text
    j = g(reg, "historias", "json").json()
    assert j["proyecto"] == C and j["vista"] == "historias" and j["borrador"] is True and {i["codigo"] for i in j["items"]} == {"HU-001", "HU-002"}
    assert [i["borrador"] for i in j["items"]] == [True, False]
    assert g(reg, "historias", "json", 0).json()["borrador"] is False
    m = g(reg, "decisiones", "md").text
    assert "BORRADOR" in m and "~~" in m and "Sustituida por:** D-" in m and "Usar Postgres" in m
    m0 = g(reg, "interesados", "md", 0).text
    assert "BORRADOR" not in m0 and "Pedro" not in m0 and "Marta Gómez" in m0


def test_vacio_con_cabeceras(reg):
    ids = [x["id"] for x in reg.get("/api/reg", params={"c": C, "tipo": "requisito"}).json()["items"]]
    reg.post("/api/reg/revisar", json={"c": C, "ids": ids, "accion": "descartar"})
    r = g(reg, "requisitos", "csv")
    assert r.status_code == 200 and r.text.lstrip("﻿").startswith("Código,Título,Detalle") and len(r.text.splitlines()) == 1
    assert g(reg, "requisitos", "docx").status_code == 200 and g(reg, "requisitos", "xlsx").status_code == 200 and g(reg, "requisitos", "md").status_code == 200


def test_docx(reg):
    for vista, cods in (("requisitos", {"REQ-001", "REQ-002"}), ("historias", {"HU-001"}), ("interesados", {"INT-001", "INT-002"}), ("decisiones", {"D-001", "D-002"})):
        r = g(reg, vista, "docx")
        assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
        d = Document(io.BytesIO(r.content))
        txt = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        assert all(c in txt for c in cods) and "BORRADOR" in txt and "descartad" not in txt.lower()
    assert "BORRADOR" not in "\n".join(p.text for p in Document(io.BytesIO(g(reg, "requisitos", "docx", 0).content)).paragraphs)
    d = Document(io.BytesIO(g(reg, "interesados", "docx").content))
    mat = next(t for t in d.tables if t.rows[0].cells[1].text == "Poder bajo")  # matriz 2x2
    assert "Marta Gómez" in mat.rows[1].cells[2].text and "Pedro Ruiz" in mat.rows[2].cells[1].text
    dec = Document(io.BytesIO(g(reg, "decisiones", "docx").content))
    assert any(r.font.strike for p in dec.paragraphs for r in p.runs) and "Sustituida por: D-002" in "\n".join(p.text for p in dec.paragraphs)


def test_xlsx(reg):
    r = g(reg, "interesados", "xlsx")
    assert r.status_code == 200 and "spreadsheetml.sheet" in r.headers["content-type"] and "interesados.xlsx" in r.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.content))
    assert z.testzip() is None
    for n in z.namelist():
        ET.fromstring(z.read(n))
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    hoja = z.read("xl/worksheets/sheet1.xml")
    filas = [[c.findtext("m:is/m:t", namespaces=ns) for c in row] for row in ET.fromstring(hoja).iter("{%s}row" % ns["m"])]
    assert filas[0][:3] == ["Código", "Nombre", "Rol"] and filas[0][-1] == "Revisión" and [f[0] for f in filas[1:]] == ["INT-001", "INT-002"]
    assert filas[2][-1].startswith("BORRADOR") and b"autoFilter" in hoja and b'state="frozen"' in hoja
    try:
        import openpyxl
    except ImportError:
        return
    assert openpyxl.load_workbook(io.BytesIO(r.content)).active["B2"].value == "Marta Gómez"


def test_feature_y_trackers(reg):
    r = g(reg, "historias", "feature")
    assert r.status_code == 200 and "historias.feature" in r.headers["content-disposition"]
    t = r.text
    assert "Feature: Reservas" in t and "Feature: Cancelar reserva" in t and t.count("Scenario:") == 3
    assert "Given un socio con sesión" in t and "When reserva una pista" in t and "Then recibe la confirmación" in t and "@propuesta" in t and "@borrador" in t
    ls = [l.strip() for l in t.splitlines() if l.strip() and not l.strip().startswith(("#", "@"))]
    assert all(l.split()[0] in ("Feature:", "Scenario:", "Given", "When", "Then", "And") for l in ls)
    t0 = g(reg, "historias", "feature", 0).text
    assert "BORRADOR" not in t0 and "Feature: Reservas" not in t0
    j = g(reg, "historias", "csv_jira")
    assert j.text.startswith("﻿Summary,Issue Type,Description,Priority,Labels") and "\r\n" in j.text and "jira.csv" in j.headers["content-disposition"]
    fj = rows(j)
    assert fj[0]["Issue Type"] == "Story" and "Scenario: Reserva válida" in fj[0]["Description"] and "borrador" in fj[0]["Labels"] and fj[0]["Story Points"] == "3"
    az = g(reg, "historias", "csv_azure")
    a = rows(az)
    assert a[0]["Work Item Type"] == "User Story" and "Given un socio con sesión" in a[0]["Acceptance Criteria"] and az.text.lstrip("﻿").startswith("Work Item Type,Title,Description")
    rq = rows(g(reg, "requisitos", "csv_jira"))
    assert [x["Issue Type"] for x in rq] == ["Epic", "Epic"] and rq[0]["Priority"] == "Highest"
    assert [x["Work Item Type"] for x in rows(g(reg, "requisitos", "csv_azure"))] == ["Requirement"] * 2


def test_pdf(reg, monkeypatch):
    with monkeypatch.context() as mp:
        mp.setattr(exportar, "navegador", lambda: None)
        r = g(reg, "requisitos", "pdf")
        assert r.status_code == 409 and "Imprimir" in r.json()["detail"]
    if os.environ.get("ACTAS_TEST_PDF") and exportar.navegador():  # lento y depende del navegador: opcional (mismo patrón que test_exportar)
        assert g(reg, "requisitos", "pdf").content.startswith(b"%PDF")
