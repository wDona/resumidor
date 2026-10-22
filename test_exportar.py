"""Pruebas de exportar.py y sus endpoints sobre una copia (acta + perfil clase) en un work temporal. `uv run pytest test_exportar.py`"""
import io
import os
import json
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

import exportar
import pipeline

SRC = Path(__file__).parent / "work" / "c8f5001c5e212d71"
VERS = ["acta_normal_rápido", "acta_normal_rápido_p-clase"]
MID = SRC.name


@pytest.fixture()
def cl(tmp_path, monkeypatch):
    w = tmp_path / "work" / MID
    w.mkdir(parents=True)
    shutil.copy(SRC / "transcript_rápido.json", w)
    for v in VERS:
        shutil.copytree(SRC / v, w / v, ignore=shutil.ignore_patterns("video.*"))
    monkeypatch.setattr(pipeline, "WORK", tmp_path / "work")
    import app
    monkeypatch.setattr(app, "CARPETAS", tmp_path / "work" / "carpetas.json")
    return TestClient(app.app)


@pytest.mark.parametrize("ver", VERS)
def test_formatos(cl, ver, monkeypatch):
    mp = monkeypatch
    g = lambda f, **q: cl.get(f"/api/exportar/{MID}/{ver}", params={"fmt": f, **q})
    from docx import Document
    r = g("docx")
    assert r.status_code == 200 and len(Document(io.BytesIO(r.content)).paragraphs) > 5
    assert g("md").text.startswith("# ")
    for f in ("obsidian", "notion"):
        z = zipfile.ZipFile(io.BytesIO(g(f).content))
        assert any(n.endswith(".md") for n in z.namelist())
        if f == "obsidian":
            assert z.read(next(n for n in z.namelist() if n.endswith(".md") and "/" not in n)).startswith(b"---")
    t = g("ics", evento=1).text
    assert t.startswith("BEGIN:VCALENDAR") and t.rstrip().endswith("END:VCALENDAR") and "BEGIN:VEVENT" in t
    assert g("csv_todoist").text.lstrip("\ufeff").startswith("TYPE,CONTENT")
    assert g("csv_jira").text.lstrip("\ufeff").startswith("Summary,")
    for f in ("srt", "vtt"):
        s = g(f, hablantes=1).text
        assert (s.startswith("WEBVTT") if f == "vtt" else s.startswith("1\n"))
    if os.environ.get("ACTAS_TEST_PDF") and exportar.navegador():  # lento y depende del navegador: opcional
        assert g("pdf").content.startswith(b"%PDF")
    mp.setattr(exportar, "navegador", lambda: None)
    r = g("pdf")
    assert r.status_code == 409 and "Imprimir" in r.json()["detail"]
    assert g("xx").status_code == 400
    assert cl.get(f"/api/vista/{MID}/{ver}?print=1").text.count("print()") == 1


def test_subtitulos():
    segs = [{"s": 0, "e": 10, "text": " ".join(["palabra"] * 30), "spk": "Ana"},
            {"s": 10, "e": 12, "text": "hola", "words": [{"w": "hola", "s": 10.5, "e": 11.0, "p": .9}], "spk": "Ana"}]
    cs = exportar.cues(segs, True)
    assert len(cs) > 2 and all(len(l) <= 42 and len(q[2].split("\n")) <= 2 for q in cs for l in q[2].split("\n"))
    assert all(a[1] <= b[0] and a[0] < a[1] for a, b in zip(cs, cs[1:])) and cs[0][2].startswith("Ana:") and "Ana:" not in cs[-1][2]
    assert cs[-1][0] == 10.5
    assert re.match(r"\d\d:\d\d:\d\d,\d{3} --> ", exportar.subtitulos(segs).split("\n")[1])


def test_edicion(cl):
    v = VERS[0]
    orig = (SRC / v / "acta.json").read_bytes()
    d = cl.get(f"/api/doc/{MID}/{v}").json()
    assert not d["editado"]
    assert cl.put(f"/api/doc/{MID}/{v}", json={"cambios": {"resumen": "RESUMEN NUEVO", "puntos.0.texto": "punto editado"}}).status_code == 200
    assert cl.put(f"/api/doc/{MID}/{v}", json={"cambios": {"nada.0.x": "a"}}).status_code == 400
    assert "RESUMEN NUEVO" in cl.get(f"/api/exportar/{MID}/{v}?fmt=md").text
    assert "RESUMEN NUEVO" not in cl.get(f"/api/vista/{MID}/{v}?original=1").text and "RESUMEN NUEVO" in cl.get(f"/api/vista/{MID}/{v}?editar=1").text
    a = next(a for a in cl.get(f"/api/meeting/{MID}").json()["actas"] if a["name"] == v)
    assert a["editado"] and a["url"].startswith("/api/vista/")
    assert (pipeline.WORK / MID / v / "acta.json").read_bytes() == orig
    cl.post(f"/api/doc/{MID}/{v}/restaurar")
    assert "RESUMEN NUEVO" not in cl.get(f"/api/exportar/{MID}/{v}?fmt=md").text


def test_estilos(cl):
    e = cl.post("/api/estilos", json={"nombre": "Mi Org", "cabecera": "Org {fecha}", "pie": "{título} p.{página}", "color": "#112233",
                                      "campos": {"lugar": "Sala 1"}, "firmas": [{"cargo": "Presidente", "nombre": "Ana"}]}).json()
    assert e["id"] == "mi-org"
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (20, 20)).save(b, "PNG")
    assert cl.put("/api/estilos/mi-org/logo?name=l.png", content=b.getvalue()).json()["logo"]
    assert cl.put("/api/estilos/mi-org/base_docx?name=b.docx", content=b"nope").status_code == 400
    cl.post("/api/carpeta", json={"nombre": "X", "estilo": "mi-org"})
    assert json.loads((pipeline.WORK / "carpetas.json").read_text())["estilos"] == {"X": "mi-org"}
    from docx import Document
    d = Document(io.BytesIO(cl.get(f"/api/exportar/{MID}/{VERS[0]}", params={"fmt": "docx", "estilo": "mi-org"}).content))
    assert "Org " in d.sections[0].header.paragraphs[0].text and "Sala 1" in "".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    cl.post("/api/estilos/borrar", json={"id": "mi-org"})
    assert cl.get("/api/estilos").json() == []
